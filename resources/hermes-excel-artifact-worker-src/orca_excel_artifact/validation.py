"""Bounded XLSX inspection, validation, and preservation-loss detection."""

from __future__ import annotations

import importlib.metadata
import os
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .archive import ZipInspection, inspect_zip, safe_read_member
from .errors import ArtifactError
from .limits import DEFAULT_LIMITS, ArtifactLimits
from .workspace import hash_regular_file


_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
_WORKBOOK_REQUIRED_PARTS = frozenset({"[Content_Types].xml", "_rels/.rels", "xl/workbook.xml"})
_SHEET_REF = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z_][^!\[\]]*))!")
_UNSUPPORTED_FEATURE_PREFIXES: dict[str, tuple[str, ...]] = {
    "activex": ("xl/activeX/",),
    "custom_xml": ("customXml/",),
    "digital_signatures": ("_xmlsignatures/",),
    "external_links": ("xl/externalLinks/",),
    "form_controls": ("xl/ctrlProps/",),
    "ole_objects": ("xl/embeddings/",),
    "pivot_caches": ("xl/pivotCache/",),
    "pivot_tables": ("xl/pivotTables/",),
    "query_tables": ("xl/queryTables/",),
    "slicer_caches": ("xl/slicerCaches/",),
    "slicers": ("xl/slicers/",),
    "threaded_comments": ("xl/threadedComments/", "xl/persons/"),
    "web_extensions": ("xl/webextensions/",),
}
_SUPPORTED_PART_PREFIXES: dict[str, tuple[str, ...]] = {
    "charts": ("xl/charts/chart",),
    "images": ("xl/media/",),
    "tables": ("xl/tables/table",),
}
_NAMESPACES = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
}
_SUPPORTED_CHART_FAMILIES = frozenset(
    {"areaChart", "barChart", "doughnutChart", "lineChart", "pieChart", "scatterChart"}
)
_CONDITIONAL_FORMAT_LOCALS = frozenset(
    {"conditionalFormatting", "cfRule", "colorScale", "dataBar", "iconSet"}
)


def _parse_xml(payload: bytes, *, part: str) -> Any:
    try:
        from defusedxml import ElementTree
    except ImportError:
        raise ArtifactError(
            "dependency_unavailable",
            "defusedxml is unavailable on the worker host.",
            "Install the pinned worker dependencies and retry.",
        ) from None
    try:
        return ElementTree.fromstring(payload)
    except Exception:
        raise ArtifactError(
            "input_format_invalid",
            "The workbook contains malformed XML.",
            "Open and re-save it with a trusted Office application.",
            {"stage": "ooxml", "feature": part[:128]},
        ) from None


def _part_count(names: frozenset[str], prefixes: tuple[str, ...]) -> int:
    lowered = tuple(prefix.casefold() for prefix in prefixes)
    return sum(any(name.casefold().startswith(prefix) for prefix in lowered) for name in names)


def _relationship_targets(
    inspection: ZipInspection,
    *,
    relationship_suffix: str,
) -> frozenset[str]:
    """Resolve OPC relationship targets of one exact relationship type."""

    targets: set[str] = set()
    for name in inspection.names:
        if not name.endswith(".rels"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        if name == "_rels/.rels":
            source_directory = ""
        else:
            rels_parent = posixpath.dirname(name)
            source_directory = posixpath.dirname(rels_parent)
        for relation in root.findall("rel:Relationship", _NAMESPACES):
            if relation.attrib.get("TargetMode") == "External":
                continue
            if not relation.attrib.get("Type", "").casefold().endswith(relationship_suffix.casefold()):
                continue
            raw_target = relation.attrib.get("Target", "").replace("\\", "/")
            resolved = posixpath.normpath(posixpath.join(source_directory, raw_target)).lstrip("/")
            if resolved and not resolved.startswith("../"):
                targets.add(resolved)
    return frozenset(targets)


def _drawing_shape_count(inspection: ZipInspection) -> int:
    count = 0
    for name in inspection.names:
        if not name.startswith("xl/drawings/drawing") or not name.endswith(".xml"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        for local_name in ("sp", "grpSp", "cxnSp", "contentPart"):
            count += len(root.findall(f".//xdr:{local_name}", _NAMESPACES))
    return count


def _tag_parts(tag: Any) -> tuple[str, str]:
    text = str(tag)
    if text.startswith("{") and "}" in text:
        namespace, local = text[1:].split("}", 1)
        return namespace, local
    return "", text


def _unsupported_chart_parts(inspection: ZipInspection) -> frozenset[str]:
    """Return charts outside the v1 native chart families or using extensions."""

    unsupported: set[str] = set()
    for name in inspection.names:
        if not name.startswith("xl/charts/chart") or not name.endswith(".xml"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        chart_families = {
            local
            for element in root.iter()
            for _, local in [_tag_parts(element.tag)]
            if local.endswith("Chart")
        }
        has_extension = any(_tag_parts(element.tag)[1] == "extLst" for element in root.iter())
        if has_extension or any(family not in _SUPPORTED_CHART_FAMILIES for family in chart_families):
            unsupported.add(name)
    return frozenset(unsupported)


def _advanced_conditional_format_parts(inspection: ZipInspection) -> frozenset[str]:
    """Find extension-based conditional formatting that openpyxl may discard."""

    advanced: set[str] = set()
    main_namespace = _NAMESPACES["main"]
    for name in inspection.names:
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        for element in root.iter():
            namespace, local = _tag_parts(element.tag)
            if local in _CONDITIONAL_FORMAT_LOCALS and namespace and namespace != main_namespace:
                advanced.add(name)
                break
        if name in advanced:
            continue
        for conditional in root.findall("main:conditionalFormatting", _NAMESPACES):
            if any(_tag_parts(element.tag)[1] == "extLst" for element in conditional.iter()):
                advanced.add(name)
                break
    return frozenset(advanced)


def _fingerprint_members(
    inspection: ZipInspection,
    names: frozenset[str] | set[str],
) -> tuple[tuple[str, int, int], ...]:
    return tuple(
        sorted(
            (member.name, member.size_bytes, member.crc32)
            for member in inspection.members
            if member.name in names
        )
    )


def _worksheet_xml_counts(inspection: ZipInspection) -> dict[str, int]:
    counts = {
        "conditional_formats": 0,
        "data_validations": 0,
        "formulas": 0,
        "merged_ranges": 0,
    }
    for name in inspection.names:
        if not name.startswith("xl/worksheets/sheet") or not name.endswith(".xml"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        counts["conditional_formats"] += len(root.findall("main:conditionalFormatting", _NAMESPACES))
        validations = root.find("main:dataValidations", _NAMESPACES)
        if validations is not None:
            counts["data_validations"] += len(validations.findall("main:dataValidation", _NAMESPACES))
        counts["formulas"] += len(root.findall(".//main:f", _NAMESPACES))
        merge_cells = root.find("main:mergeCells", _NAMESPACES)
        if merge_cells is not None:
            counts["merged_ranges"] += len(merge_cells.findall("main:mergeCell", _NAMESPACES))
    return counts


@dataclass(frozen=True, slots=True)
class FeatureInventory:
    """Counts stable enough to compare before and after an openpyxl save."""

    counts: dict[str, int]
    parts: frozenset[str] = field(repr=False)
    preservation_fingerprints: dict[str, tuple[tuple[str, int, int], ...]] = field(
        repr=False,
        default_factory=dict,
    )

    def public_dict(self) -> dict[str, int]:
        return dict(sorted(self.counts.items()))


@dataclass(frozen=True, slots=True)
class PreservationReport:
    lost: dict[str, int]
    warnings: tuple[dict[str, Any], ...]

    @property
    def passed(self) -> bool:
        return not self.lost


@dataclass(slots=True)
class ValidationReport:
    passed: bool
    checks: list[dict[str, str]]
    inventory: FeatureInventory
    sheets: list[dict[str, Any]]
    tables: list[dict[str, Any]]
    charts: list[dict[str, Any]]
    warnings: list[dict[str, Any]] = field(default_factory=list)
    render_pdf: Path | None = field(default=None, repr=False)

    def public_dict(self) -> dict[str, Any]:
        return {"passed": self.passed, "checks": self.checks}


def inspect_ooxml(path: Path | str, *, limits: ArtifactLimits = DEFAULT_LIMITS) -> tuple[ZipInspection, FeatureInventory]:
    inspection = inspect_zip(path, limits=limits)
    missing = sorted(_WORKBOOK_REQUIRED_PARTS - inspection.names)
    if missing:
        raise ArtifactError(
            "input_format_invalid",
            "The XLSX package is missing a required part.",
            "Open and re-save it as a standard .xlsx workbook.",
            {"stage": "ooxml", "feature": missing[0][:128]},
        )
    content_types = _parse_xml(
        safe_read_member(inspection, "[Content_Types].xml", max_bytes=4 * 1024 * 1024),
        part="[Content_Types].xml",
    )
    content_type_values = {
        node.attrib.get("ContentType", "")
        for node in content_types.findall("ct:Override", _NAMESPACES)
    }
    if _XLSX_CONTENT_TYPE not in content_type_values:
        raise ArtifactError(
            "input_format_invalid",
            "The package is not a standard editable XLSX workbook.",
            "Use a non-encrypted .xlsx workbook without a renamed extension.",
        )
    workbook_root = _parse_xml(
        safe_read_member(inspection, "xl/workbook.xml", max_bytes=8 * 1024 * 1024),
        part="xl/workbook.xml",
    )
    sheet_nodes = workbook_root.findall("main:sheets/main:sheet", _NAMESPACES)
    defined_names = workbook_root.find("main:definedNames", _NAMESPACES)
    counts: dict[str, int] = {
        "defined_names": len(defined_names) if defined_names is not None else 0,
        "drawing_shapes": _drawing_shape_count(inspection),
        "sheets": len(sheet_nodes),
    }
    counts.update(_worksheet_xml_counts(inspection))
    for feature, prefixes in _SUPPORTED_PART_PREFIXES.items():
        counts[feature] = _part_count(inspection.names, prefixes)
    for feature, prefixes in _UNSUPPORTED_FEATURE_PREFIXES.items():
        counts[feature] = _part_count(inspection.names, prefixes)
    ole_parts = _relationship_targets(inspection, relationship_suffix="/oleObject")
    # `xl/embeddings/*.xlsx` commonly backs a native chart.  Only parts reached
    # by an explicit oleObject relationship are OLE objects.
    counts["ole_objects"] = len(ole_parts)
    counts["macros"] = int("xl/vbaProject.bin" in inspection.names)
    unsupported_chart_parts = _unsupported_chart_parts(inspection)
    advanced_conditional_parts = _advanced_conditional_format_parts(inspection)
    theme_parts = frozenset(
        name for name in inspection.names if name.startswith("xl/theme/") and name.endswith(".xml")
    )
    counts["unsupported_charts"] = len(unsupported_chart_parts)
    counts["advanced_conditional_formats"] = len(advanced_conditional_parts)
    counts["theme_parts"] = len(theme_parts)
    fingerprints = {
        feature: tuple(
            sorted(
                (member.name, member.size_bytes, member.crc32)
                for member in inspection.members
                if any(member.name.startswith(prefix) for prefix in prefixes)
            )
        )
        for feature, prefixes in _UNSUPPORTED_FEATURE_PREFIXES.items()
    }
    fingerprints["ole_objects"] = tuple(
        sorted(
            (member.name, member.size_bytes, member.crc32)
            for member in inspection.members
            if member.name in ole_parts
        )
    )
    fingerprints["unsupported_charts"] = _fingerprint_members(inspection, unsupported_chart_parts)
    fingerprints["advanced_conditional_formats"] = _fingerprint_members(
        inspection, advanced_conditional_parts
    )
    # Themes are usually safe to retain and therefore are not a preflight
    # blocker, but a byte-level change or loss after an openpyxl round trip is
    # material and must be reported.
    fingerprints["theme_parts"] = _fingerprint_members(inspection, theme_parts)
    return inspection, FeatureInventory(counts, inspection.names, fingerprints)


def compare_feature_inventories(
    before: FeatureInventory,
    after: FeatureInventory,
    *,
    preservation_policy: str,
) -> PreservationReport:
    """Detect actual feature count loss and enforce the requested policy."""

    lost = {
        feature: count - after.counts.get(feature, 0)
        for feature, count in before.counts.items()
        if count > after.counts.get(feature, 0)
    }
    for feature, before_fingerprint in before.preservation_fingerprints.items():
        if before_fingerprint and before_fingerprint != after.preservation_fingerprints.get(feature, ()):
            lost[feature] = max(lost.get(feature, 0), len(before_fingerprint))
    warnings = tuple(
        {
            "code": "feature_loss_detected",
            "message": f"Modification lost {count} instance(s) of workbook feature '{feature}'.",
            "details": {"stage": "preservation", "feature": feature, "count": count},
        }
        for feature, count in sorted(lost.items())
    )
    if lost and preservation_policy == "fail_on_unsupported_loss":
        raise ArtifactError(
            "validation_failed",
            "The modified workbook lost features while being re-serialized.",
            "Use a supported workbook, remove the lossy edit, or explicitly choose warn_on_unsupported_loss.",
            {"stage": "preservation", "feature": next(iter(sorted(lost)))},
        )
    return PreservationReport(lost, warnings)


def preflight_unsupported_features(
    inventory: FeatureInventory,
    *,
    preservation_policy: str,
) -> tuple[dict[str, Any], ...]:
    """Fail or warn before openpyxl sees known round-trip-unsafe features."""

    present = {
        feature: inventory.counts.get(feature, 0)
        for feature in (
            *_UNSUPPORTED_FEATURE_PREFIXES,
            "macros",
            "drawing_shapes",
            "unsupported_charts",
            "advanced_conditional_formats",
        )
        if inventory.counts.get(feature, 0) > 0
    }
    warnings = tuple(
        {
            "code": "unsupported_feature_preservation_risk",
            "message": f"The source contains {count} instance(s) of '{feature}', which openpyxl may not preserve.",
            "details": {"stage": "preservation", "feature": feature, "count": count},
        }
        for feature, count in sorted(present.items())
    )
    if present and preservation_policy == "fail_on_unsupported_loss":
        raise ArtifactError(
            "capability_unsupported",
            "The source workbook contains a feature that cannot be safely round-tripped by openpyxl.",
            "Use a workbook without the feature or explicitly choose warn_on_unsupported_loss.",
            {"stage": "preservation", "feature": next(iter(sorted(present)))},
        )
    return warnings


def _formula_sheet_names(formula: str) -> list[str]:
    return [
        (quoted.replace("''", "'") if quoted else unquoted.strip())
        for quoted, unquoted in _SHEET_REF.findall(formula)
    ]


def _validate_formula_references(workbook: Any) -> tuple[bool, str]:
    sheet_names = {name.casefold() for name in workbook.sheetnames}
    checked = 0
    invalid: list[str] = []
    try:
        from openpyxl.formula import Tokenizer
    except ImportError:
        return False, "openpyxl formula tokenizer is unavailable."
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                checked += 1
                if "#REF!" in value.upper():
                    invalid.append(f"{worksheet.title}!{cell.coordinate}")
                    continue
                try:
                    Tokenizer(value)
                except Exception:
                    invalid.append(f"{worksheet.title}!{cell.coordinate}")
                    continue
                if any(name.casefold() not in sheet_names for name in _formula_sheet_names(value)):
                    invalid.append(f"{worksheet.title}!{cell.coordinate}")
    if invalid:
        return False, f"{len(invalid)} formula(s) contain invalid or missing-sheet references."
    return True, f"Parsed {checked} formula(s) without broken sheet references."


def _validate_defined_name_references(workbook: Any) -> tuple[bool, str]:
    """Validate defined-name scope and worksheet/range destinations."""

    try:
        from openpyxl.utils.cell import range_boundaries
    except ImportError:
        return False, "openpyxl range validation is unavailable."
    sheet_names = {name.casefold() for name in workbook.sheetnames}
    invalid = 0
    checked = 0
    for defined_name in workbook.defined_names.values():
        checked += 1
        text = str(defined_name.attr_text or "")
        if not text or "#REF!" in text.upper() or "[" in text or "]" in text:
            invalid += 1
            continue
        local_sheet_id = getattr(defined_name, "localSheetId", None)
        if local_sheet_id is not None and not (0 <= int(local_sheet_id) < len(workbook.sheetnames)):
            invalid += 1
            continue
        referenced_sheets = _formula_sheet_names(text)
        if any(name.casefold() not in sheet_names for name in referenced_sheets):
            invalid += 1
            continue
        if "!" not in text:
            # Constant/formula defined names are valid but do not expose a
            # direct worksheet range for bounds checking.
            continue
        try:
            destinations = list(defined_name.destinations)
        except Exception:
            invalid += 1
            continue
        if not destinations:
            invalid += 1
            continue
        try:
            for sheet_name, coordinate in destinations:
                if sheet_name.casefold() not in sheet_names:
                    raise ValueError("missing sheet")
                range_boundaries(str(coordinate).replace("$", ""))
        except (TypeError, ValueError):
            invalid += 1
    if invalid:
        return False, f"{invalid} defined name(s) contain an invalid scope, sheet, or range reference."
    return True, f"Checked {checked} defined name reference(s)."


def _chart_formulas(inspection: ZipInspection) -> list[str]:
    formulas: list[str] = []
    for name in inspection.names:
        if not name.startswith("xl/charts/chart") or not name.endswith(".xml"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        formulas.extend(node.text or "" for node in root.findall(".//c:f", _NAMESPACES))
    return formulas


def _validate_chart_references(inspection: ZipInspection, sheet_names: set[str]) -> tuple[bool, str]:
    formulas = _chart_formulas(inspection)
    invalid = 0
    for formula in formulas:
        if "#REF!" in formula.upper() or any(
            name.casefold() not in sheet_names for name in _formula_sheet_names(formula)
        ):
            invalid += 1
    if invalid:
        return False, f"{invalid} chart reference(s) are broken or target a missing sheet."
    return True, f"Checked {len(formulas)} chart reference formula(s)."


def _engine_version(name: str) -> str:
    try:
        return importlib.metadata.version(name)
    except importlib.metadata.PackageNotFoundError:
        return "unavailable"


def _openpyxl_inspect(path: Path) -> tuple[Any, list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    try:
        import openpyxl
        workbook = openpyxl.load_workbook(path, read_only=False, data_only=False, keep_links=True)
    except ImportError:
        raise ArtifactError(
            "dependency_unavailable",
            "openpyxl is unavailable on the worker host.",
            "Install the pinned worker dependencies and retry.",
        ) from None
    except Exception:
        raise ArtifactError(
            "validation_failed",
            "openpyxl could not reopen the generated workbook.",
            "Correct workbook corruption or unsupported package features and retry.",
            {"stage": "openpyxl"},
        ) from None
    sheets = [
        {"name": sheet.title, "status": "validated"}
        for sheet in workbook.worksheets
    ]
    tables = [
        {"name": table.name, "sheet": sheet.title, "range": table.ref, "status": "validated"}
        for sheet in workbook.worksheets
        for table in sheet.tables.values()
    ]
    charts = [
        {"name": f"chart-{index}", "sheet": sheet.title, "status": "validated"}
        for sheet in workbook.worksheets
        for index, _ in enumerate(sheet._charts, start=1)
    ]
    return workbook, sheets, tables, charts


def _check(name: str, passed: bool, success: str, failure: str) -> dict[str, str]:
    return {
        "name": name,
        "status": "passed" if passed else "failed",
        "message": success if passed else failure,
    }


def _validated_spec(spec: Mapping[str, Any], action: str) -> Mapping[str, Any]:
    # Local import keeps the validation module usable by the workbook engines
    # without introducing a module-import cycle.
    from .workbooks import validate_workbook_spec

    return validate_workbook_spec(spec, action=action).spec


def _normal_formula_for_comparison(value: Any) -> str:
    text = str(value)
    return text if text.startswith("=") else f"={text}"


def _normal_range(value: Any) -> str:
    return str(value).replace("$", "").upper()


def _normal_defined_name_formula(value: Any) -> str:
    text = str(value or "").strip()
    return text[1:] if text.startswith("=") else text


def _normal_print_area(value: Any) -> str:
    if value is None:
        return ""
    text = str(getattr(value, "attr_text", value)).strip()
    if "!" in text:
        text = text.rsplit("!", 1)[1]
    return _normal_range(text)


def _literal_matches(actual: Any, expected: Any) -> bool:
    if expected is None:
        return actual is None
    if isinstance(expected, bool):
        return isinstance(actual, bool) and actual is expected
    if isinstance(expected, (int, float)) and not isinstance(expected, bool):
        return isinstance(actual, (int, float)) and not isinstance(actual, bool) and actual == expected
    return type(actual) is type(expected) and actual == expected


def _expected_cells(sheet_spec: Mapping[str, Any]) -> dict[str, tuple[str, Any]]:
    """Return final declared cell content after explicit-cell precedence."""

    try:
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ArtifactError(
            "dependency_unavailable",
            "openpyxl is unavailable on the worker host.",
            "Install the pinned worker dependencies and retry.",
        ) from None
    expected: dict[str, tuple[str, Any]] = {}
    for row_index, row in enumerate(sheet_spec.get("data", []), start=1):
        for column_index, item in enumerate(row, start=1):
            address = f"{get_column_letter(column_index)}{row_index}"
            if isinstance(item, Mapping) and "formula" in item:
                expected[address] = ("formula", _normal_formula_for_comparison(item["formula"]))
            elif isinstance(item, Mapping):
                expected[address] = ("literal", item.get("value"))
            else:
                expected[address] = ("literal", item)
    for item in sheet_spec.get("cells", []):
        address = str(item["address"]).replace("$", "").upper()
        if "formula" in item:
            expected[address] = ("formula", _normal_formula_for_comparison(item["formula"]))
        else:
            expected[address] = ("literal", item.get("value"))
    return expected


def _normal_chart_reference(value: Any) -> str:
    text = str(value or "").strip().lstrip("=")
    if "!" not in text:
        return _normal_range(text)
    sheet, cell_range = text.rsplit("!", 1)
    sheet = sheet.strip()
    if len(sheet) >= 2 and sheet[0] == sheet[-1] == "'":
        sheet = sheet[1:-1].replace("''", "'")
    return f"{sheet.casefold()}!{_normal_range(cell_range)}"


def _chart_xml_signatures(
    inspection: ZipInspection,
) -> list[tuple[tuple[str, ...], tuple[str, ...]]]:
    signatures: list[tuple[tuple[str, ...], tuple[str, ...]]] = []
    for name in sorted(inspection.names):
        if not name.startswith("xl/charts/chart") or not name.endswith(".xml"):
            continue
        root = _parse_xml(safe_read_member(inspection, name), part=name)
        categories: list[str] = []
        values: list[str] = []
        for series in root.findall(".//c:ser", _NAMESPACES):
            category_nodes = series.findall(".//c:cat//c:f", _NAMESPACES)
            category_nodes += series.findall(".//c:xVal//c:f", _NAMESPACES)
            value_nodes = series.findall(".//c:val//c:f", _NAMESPACES)
            value_nodes += series.findall(".//c:yVal//c:f", _NAMESPACES)
            categories.extend(_normal_chart_reference(node.text) for node in category_nodes)
            values.extend(_normal_chart_reference(node.text) for node in value_nodes)
        signatures.append((tuple(sorted(categories)), tuple(sorted(values))))
    return signatures


def _expected_chart_signatures(
    spec: Mapping[str, Any],
) -> list[tuple[tuple[str, ...], tuple[str, ...]]]:
    signatures: list[tuple[tuple[str, ...], tuple[str, ...]]] = []
    for sheet_spec in spec.get("sheets", []):
        for chart in sheet_spec.get("charts", []):
            categories = [
                _normal_chart_reference(
                    f"{series['categories']['sheet']}!{series['categories']['range']}"
                )
                for series in chart.get("series", [])
                if "categories" in series
            ]
            values = [
                _normal_chart_reference(f"{series['values']['sheet']}!{series['values']['range']}")
                for series in chart.get("series", [])
            ]
            signatures.append((tuple(sorted(categories)), tuple(sorted(values))))
    return signatures


def _contains_chart_signatures(
    actual: list[tuple[tuple[str, ...], tuple[str, ...]]],
    expected: list[tuple[tuple[str, ...], tuple[str, ...]]],
) -> bool:
    remaining = Counter(actual)
    for signature in expected:
        if remaining[signature] <= 0:
            return False
        remaining[signature] -= 1
    return True


def _validate_spec_with_open_workbook(
    workbook: Any,
    inspection: ZipInspection,
    spec: Mapping[str, Any],
    action: str,
) -> list[dict[str, str]]:
    if action not in {"create", "modify"}:
        raise ArtifactError(
            "protocol_invalid",
            "Workbook spec validation requires action create or modify.",
            "Use the action that produced the workbook.",
            {"field": "action"},
        )

    checks: list[dict[str, str]] = []
    declared_sheets = list(spec.get("sheets", []))
    expected_names = [str(item.get("name", "")) for item in declared_sheets]
    actual_names = list(workbook.sheetnames)
    declared_sheets_present = all(name in workbook.sheetnames for name in expected_names)
    if action == "create":
        sheets_ok = actual_names == expected_names
        sheet_success = f"Matched the exact declared order of {len(expected_names)} sheet(s)."
        sheet_failure = "The created workbook's sheet names or order differ from the declaration."
    else:
        sheets_ok = declared_sheets_present
        sheet_success = f"Found all {len(expected_names)} declared sheet(s) in the modified workbook."
        sheet_failure = "The modified workbook is missing one or more declared sheets."
    if sheets_ok:
        sheets_ok = all(
            workbook[item["name"]].sheet_state == item.get("state")
            for item in declared_sheets
        )
        if not sheets_ok:
            sheet_failure = "One or more declared sheets have the wrong visibility state."
    checks.append(_check("spec.sheets", sheets_ok, sheet_success, sheet_failure))

    cell_total = 0
    cell_failures = 0
    for sheet_spec in declared_sheets:
        if sheet_spec.get("name") not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_spec["name"]]
        for address, (kind, expected) in _expected_cells(sheet_spec).items():
            cell_total += 1
            cell = worksheet[address]
            if kind == "formula":
                matches = cell.data_type == "f" and cell.value == expected
            else:
                matches = cell.data_type != "f" and _literal_matches(cell.value, expected)
            cell_failures += int(not matches)
    checks.append(_check(
        "spec.cells",
        cell_failures == 0 and declared_sheets_present,
        f"Matched {cell_total} declared formula and literal cell value(s).",
        f"{cell_failures or 1} declared cell value(s) are missing or differ.",
    ))

    merge_total = 0
    merge_failures = 0
    freeze_total = 0
    freeze_failures = 0
    filter_total = 0
    filter_failures = 0
    print_total = 0
    print_failures = 0
    for sheet_spec in declared_sheets:
        if sheet_spec.get("name") not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_spec["name"]]
        expected_merges = {_normal_range(item) for item in sheet_spec.get("merges", [])}
        actual_merges = {_normal_range(item) for item in worksheet.merged_cells.ranges}
        merge_total += len(expected_merges)
        if action == "create":
            merge_failures += len(expected_merges.symmetric_difference(actual_merges))
        else:
            merge_failures += len(expected_merges - actual_merges)

        if "freezePane" in sheet_spec:
            freeze_total += 1
            freeze = sheet_spec["freezePane"]
            row = int(freeze.get("row", 0))
            column = int(freeze.get("column", 0))
            expected_freeze = None if row == column == 0 else worksheet.cell(row + 1, column + 1).coordinate
            actual_freeze = getattr(worksheet.freeze_panes, "coordinate", worksheet.freeze_panes)
            freeze_failures += int(actual_freeze != expected_freeze)

        if sheet_spec.get("autofilter"):
            filter_total += 1
            expected_filter = _normal_range(sheet_spec["autofilter"]["range"])
            actual_filter = _normal_range(worksheet.auto_filter.ref or "")
            table_filters = {_normal_range(table.ref) for table in worksheet.tables.values()}
            filter_failures += int(expected_filter != actual_filter and expected_filter not in table_filters)

        page_setup = sheet_spec.get("pageSetup", {})
        if "printArea" in page_setup:
            print_total += 1
            print_failures += int(
                _normal_print_area(worksheet.print_area) != _normal_range(page_setup["printArea"])
            )
    checks.extend([
        _check("spec.merges", merge_failures == 0 and declared_sheets_present,
               f"Matched {merge_total} declared merged range(s).",
               "One or more declared merged ranges are missing or differ."),
        _check("spec.freeze_panes", freeze_failures == 0 and declared_sheets_present,
               f"Matched {freeze_total} declared freeze-pane setting(s).",
               "One or more declared freeze-pane settings are missing or differ."),
        _check("spec.autofilters", filter_failures == 0 and declared_sheets_present,
               f"Matched {filter_total} declared autofilter range(s).",
               "One or more declared autofilter ranges are missing or differ."),
        _check("spec.print_areas", print_failures == 0 and declared_sheets_present,
               f"Matched {print_total} declared print area(s).",
               "One or more declared print areas are missing or differ."),
    ])

    expected_names_formulas = {
        str(item["name"]).casefold(): _normal_defined_name_formula(item["formula"])
        for item in spec.get("namedRanges", [])
    }
    actual_names_formulas = {
        str(name).casefold(): _normal_defined_name_formula(value.attr_text)
        for name, value in workbook.defined_names.items()
    }
    names_ok = all(
        actual_names_formulas.get(name) == formula
        for name, formula in expected_names_formulas.items()
    )
    if action == "create":
        names_ok = names_ok and set(actual_names_formulas) == set(expected_names_formulas)
    checks.append(_check(
        "spec.named_ranges",
        names_ok,
        f"Matched {len(expected_names_formulas)} declared named range(s).",
        "One or more declared named ranges are missing or differ.",
    ))

    expected_tables = {
        (str(table["name"]), str(sheet["name"]), _normal_range(table["range"]))
        for sheet in declared_sheets
        for table in sheet.get("tables", [])
    }
    actual_tables = {
        (str(table.name), str(worksheet.title), _normal_range(table.ref))
        for worksheet in workbook.worksheets
        for table in worksheet.tables.values()
    }
    tables_ok = expected_tables == actual_tables if action == "create" else expected_tables <= actual_tables
    checks.append(_check(
        "spec.tables",
        tables_ok,
        f"Matched {len(expected_tables)} declared native table(s) by name and range.",
        "One or more declared native table names or ranges are missing or differ.",
    ))

    expected_charts = _expected_chart_signatures(spec)
    actual_charts = _chart_xml_signatures(inspection)
    chart_count_ok = (
        len(actual_charts) == len(expected_charts)
        if action == "create"
        else len(actual_charts) >= len(expected_charts)
    )
    for sheet_spec in declared_sheets:
        if sheet_spec.get("name") not in workbook.sheetnames:
            chart_count_ok = False
            continue
        expected_sheet_count = len(sheet_spec.get("charts", []))
        actual_sheet_count = len(workbook[sheet_spec["name"]]._charts)
        if action == "create":
            chart_count_ok = chart_count_ok and actual_sheet_count == expected_sheet_count
        else:
            chart_count_ok = chart_count_ok and actual_sheet_count >= expected_sheet_count
    charts_ok = chart_count_ok and _contains_chart_signatures(actual_charts, expected_charts)
    checks.append(_check(
        "spec.charts",
        charts_ok,
        f"Matched {len(expected_charts)} declared native chart(s) and their series references.",
        "The native chart count or one or more declared category/value references differ.",
    ))
    return checks


def validate_against_spec(
    path: Path | str,
    spec: Mapping[str, Any],
    action: str,
    *,
    limits: ArtifactLimits = DEFAULT_LIMITS,
) -> list[dict[str, str]]:
    """Validate a generated XLSX against Workbook Spec v1 without exposing its path."""

    validated_spec = _validated_spec(spec, action)
    workbook_path = Path(path)
    inspection, _ = inspect_ooxml(workbook_path, limits=limits)
    workbook, _, _, _ = _openpyxl_inspect(workbook_path)
    try:
        return _validate_spec_with_open_workbook(workbook, inspection, validated_spec, action)
    finally:
        workbook.close()


def find_libreoffice() -> Path | None:
    for name in ("libreoffice", "soffice"):
        found = shutil.which(name)
        if found:
            return Path(found)
    if os.name == "nt":
        for candidate in (
            Path(os.environ.get("ProgramFiles", "")) / "LibreOffice/program/soffice.exe",
            Path(os.environ.get("ProgramFiles(x86)", "")) / "LibreOffice/program/soffice.exe",
        ):
            if candidate.is_file():
                return candidate
    return None


def render_workbook_pdf(
    path: Path | str,
    output_dir: Path | str,
    *,
    timeout_seconds: int = 120,
    executable: Path | None = None,
) -> Path:
    """Render with isolated LibreOffice state; never invoke a shell."""

    source = Path(path)
    destination = Path(output_dir)
    destination.mkdir(mode=0o700, parents=True, exist_ok=True)
    office = executable or find_libreoffice()
    if office is None:
        raise ArtifactError(
            "render_unavailable",
            "LibreOffice rendering is unavailable on the worker host.",
            "Install LibreOffice or submit without renderPreview validation.",
        )
    with tempfile.TemporaryDirectory(prefix="orca-lo-profile-") as profile:
        profile_uri = Path(profile).resolve().as_uri()
        command = [
            str(office),
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            f"-env:UserInstallation={profile_uri}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(destination),
            str(source),
        ]
        try:
            completed = subprocess.run(
                command,
                stdin=subprocess.DEVNULL,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                timeout=timeout_seconds,
                check=False,
                shell=False,
                env={**os.environ, "HOME": profile},
            )
        except subprocess.TimeoutExpired:
            raise ArtifactError(
                "timeout",
                "LibreOffice workbook rendering timed out.",
                "Use a smaller workbook or retry with a reviewed timeout.",
                {"stage": "render"},
            ) from None
        except OSError:
            raise ArtifactError(
                "render_failed",
                "LibreOffice could not be started.",
                "Repair the worker's LibreOffice installation and retry.",
                {"stage": "render"},
            ) from None
        target = destination / f"{source.stem}.pdf"
        if completed.returncode != 0 or not target.is_file() or target.stat().st_size == 0:
            raise ArtifactError(
                "render_failed",
                "LibreOffice did not produce a usable PDF preview.",
                "Open the workbook manually or correct unsupported render content.",
                {"stage": "render"},
            )
        return target


def validate_workbook(
    path: Path | str,
    request: dict[str, Any] | None = None,
    *,
    limits: ArtifactLimits = DEFAULT_LIMITS,
    render_dir: Path | None = None,
    timeout_seconds: int = 120,
    workbook_spec: Mapping[str, Any] | None = None,
    action: str | None = None,
) -> ValidationReport:
    """Run structural, editable-object, reference, optional spec, and render checks."""

    if workbook_spec is not None and action not in {"create", "modify"}:
        raise ArtifactError(
            "protocol_invalid",
            "Workbook spec validation requires action create or modify.",
            "Use the action that produced the workbook.",
            {"field": "action"},
        )
    validated_spec = _validated_spec(workbook_spec, action or "") if workbook_spec is not None else None
    options = request or {"openXml": True, "formulas": True, "charts": True, "renderPreview": False}
    checks: list[dict[str, str]] = []
    workbook_path = Path(path)
    inspection, inventory = inspect_ooxml(workbook_path, limits=limits)
    checks.append({"name": "ooxml.package", "status": "passed", "message": "OOXML package structure and XML parsed successfully."})
    workbook, sheets, tables, charts = _openpyxl_inspect(workbook_path)
    checks.append({"name": "openpyxl.reopen", "status": "passed", "message": f"openpyxl {_engine_version('openpyxl')} reopened the editable workbook."})
    if options.get("formulas", True):
        passed, message = _validate_formula_references(workbook)
        checks.append({"name": "formula.references", "status": "passed" if passed else "failed", "message": message})
        passed, message = _validate_defined_name_references(workbook)
        checks.append({"name": "defined_name.references", "status": "passed" if passed else "failed", "message": message})
    else:
        checks.append({"name": "formula.references", "status": "skipped", "message": "Formula reference validation was not requested."})
        checks.append({"name": "defined_name.references", "status": "skipped", "message": "Defined-name reference validation was not requested."})
    if options.get("charts", True):
        passed, message = _validate_chart_references(inspection, {name.casefold() for name in workbook.sheetnames})
        checks.append({"name": "chart.references", "status": "passed" if passed else "failed", "message": message})
    else:
        checks.append({"name": "chart.references", "status": "skipped", "message": "Chart reference validation was not requested."})
    if validated_spec is not None:
        checks.extend(
            _validate_spec_with_open_workbook(
                workbook,
                inspection,
                validated_spec,
                action or "",
            )
        )
    render_pdf: Path | None = None
    if options.get("renderPreview", False):
        try:
            render_pdf = render_workbook_pdf(
                workbook_path,
                render_dir or workbook_path.parent,
                timeout_seconds=timeout_seconds,
            )
            checks.append({"name": "libreoffice.render", "status": "passed", "message": "LibreOffice produced a non-empty PDF preview."})
        except ArtifactError as exc:
            checks.append({
                "name": "libreoffice.render",
                "status": "unavailable" if exc.code == "render_unavailable" else "failed",
                "message": exc.message,
            })
    else:
        checks.append({"name": "libreoffice.render", "status": "skipped", "message": "Optional LibreOffice rendering was not requested."})
    workbook.close()
    passed = not any(check["status"] in {"failed", "unavailable"} for check in checks)
    return ValidationReport(passed, checks, inventory, sheets, tables, charts, render_pdf=render_pdf)


def inspect_xlsx(
    path: Path | str,
    *,
    expected_sha256: str | None = None,
    limits: ArtifactLimits = DEFAULT_LIMITS,
) -> dict[str, Any]:
    """Public, path-free XLSX inspection used by the generic input router."""

    workbook_path = Path(path)
    digest, size = hash_regular_file(workbook_path, max_bytes=limits.max_input_bytes)
    if expected_sha256 is not None and digest != expected_sha256.casefold():
        raise ArtifactError(
            "input_hash_mismatch",
            "The workbook input hash does not match the admitted artifact.",
            "Attach the current workbook and retry.",
            {"expectedSha256": expected_sha256.casefold(), "actualSha256": digest},
        )
    report = validate_workbook(
        workbook_path,
        {"openXml": True, "formulas": False, "charts": False, "renderPreview": False},
        limits=limits,
    )
    return {
        "kind": "xlsx",
        "sha256": digest,
        "sizeBytes": size,
        "sheetCount": len(report.sheets),
        "sheets": [item["name"] for item in report.sheets],
        "tableCount": len(report.tables),
        "chartCount": len(report.charts),
        "features": report.inventory.public_dict(),
    }
