"""Declarative Workbook Spec v1 execution engines.

Creation uses XlsxWriter because it produces native editable Office objects.
Modification uses openpyxl and always operates on a private staged copy.
"""

from __future__ import annotations

import copy
import re
from collections.abc import Callable, Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .errors import ArtifactError


WORKBOOK_SPEC_VERSION = 1
PRESERVATION_POLICIES = frozenset({"fail_on_unsupported_loss", "warn_on_unsupported_loss"})
SHEET_STATES = frozenset({"visible", "hidden", "veryHidden"})
_CELL = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]{0,6})$")
_CELL_RANGE = re.compile(
    r"^\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6}(?::\$?[A-Z]{1,3}\$?[1-9][0-9]{0,6})?$"
)
_COLUMN_RANGE = re.compile(r"^[A-Z]{1,3}(?::[A-Z]{1,3})?$")
_SHEET_FORBIDDEN = frozenset(":*?/\\[]")
_FORMAT_KEYS = frozenset({"id", "font", "fill", "alignment", "border", "numberFormat", "locked"})
_SHEET_KEYS = frozenset(
    {
        "name",
        "state",
        "tabColor",
        "data",
        "cells",
        "columns",
        "rows",
        "merges",
        "freezePane",
        "autofilter",
        "tables",
        "charts",
        "dataValidations",
        "conditionalFormats",
        "images",
        "pageSetup",
    }
)
_SPEC_KEYS = frozenset({"version", "preservationPolicy", "properties", "formats", "namedRanges", "sheets"})
_CELL_KEYS = frozenset({"address", "value", "formula", "formatId", "numberFormat", "comment", "hyperlink"})
_XLSXWRITER_CHART_TYPES = frozenset({"area", "bar", "column", "doughnut", "line", "pie", "scatter"})
_FORBIDDEN_FORMULA_FUNCTION = re.compile(
    r"(?i)(?:^|[^A-Z0-9_.])(?:_XLFN\.)?"
    r"(?:WEBSERVICE|FILTERXML|RTD|CALL|REGISTER\.ID|EXEC)\s*\("
)
_EXTERNAL_WORKBOOK_REFERENCE = re.compile(r"(?i)\[[^\]\r\n]+\][^!\r\n]*!")
_NETWORK_REFERENCE = re.compile(r"(?i)(?:https?|ftps?|file|smb)://|\\\\")


def _unsupported(message: str, *, field_path: str | None = None) -> ArtifactError:
    details = {"field": field_path} if field_path else {}
    return ArtifactError(
        "capability_unsupported",
        message,
        "Remove the unsupported Workbook Spec v1 field or use an advertised future capability.",
        details,
    )


def _protocol(message: str, *, field_path: str | None = None) -> ArtifactError:
    details = {"field": field_path} if field_path else {}
    return ArtifactError(
        "protocol_invalid",
        message,
        "Correct the declarative Workbook Spec v1 request and retry.",
        details,
    )


def _strict_keys(value: Mapping[str, Any], allowed: frozenset[str], path: str) -> None:
    unknown = sorted(set(value) - allowed)
    if unknown:
        raise _unsupported(f"Workbook Spec v1 does not support field '{unknown[0]}'.", field_path=f"{path}.{unknown[0]}")


def _require_mapping(value: Any, path: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise _protocol(f"{path} must be an object.", field_path=path)
    return value


def _require_list(value: Any, path: str) -> list[Any]:
    if not isinstance(value, list):
        raise _protocol(f"{path} must be an array.", field_path=path)
    return value


def _column_index(label: str) -> int:
    result = 0
    for character in label:
        result = result * 26 + ord(character) - 64
    return result


def _validate_address(value: Any, path: str) -> str:
    if not isinstance(value, str) or not (match := _CELL.fullmatch(value)):
        raise _protocol(f"{path} must be an A1 cell address.", field_path=path)
    if int(match.group(2)) > 1_048_576 or _column_index(match.group(1)) > 16_384:
        raise _protocol(f"{path} exceeds Excel worksheet bounds.", field_path=path)
    return value


def _validate_range(value: Any, path: str) -> str:
    if not isinstance(value, str) or not _CELL_RANGE.fullmatch(value):
        raise _protocol(f"{path} must be an A1 cell range.", field_path=path)
    for address in value.split(":"):
        _validate_address(address, path)
    return value


def _normal_formula(value: Any, path: str) -> str:
    if not isinstance(value, str) or not value or len(value) > 8192:
        raise _protocol(f"{path} must be a non-empty formula.", field_path=path)
    formula = value if value.startswith("=") else f"={value}"
    if _FORBIDDEN_FORMULA_FUNCTION.search(formula):
        raise _unsupported("The formula uses a network or execution-capable function.", field_path=path)
    if _EXTERNAL_WORKBOOK_REFERENCE.search(formula):
        raise _unsupported("External-workbook formula references are forbidden in Workbook Spec v1.", field_path=path)
    if "|" in formula:
        raise _unsupported("DDE-style formula references are forbidden in Workbook Spec v1.", field_path=path)
    if _NETWORK_REFERENCE.search(formula):
        raise _unsupported("Network references are forbidden in Workbook Spec v1 formulas.", field_path=path)
    return formula


@dataclass(frozen=True, slots=True)
class WorkbookPlan:
    spec: dict[str, Any]
    preservation_policy: str
    formats: dict[str, dict[str, Any]]
    sheet_names: tuple[str, ...]
    warnings: tuple[dict[str, Any], ...] = ()


def validate_workbook_spec(spec: Mapping[str, Any], *, action: str) -> WorkbookPlan:
    """Strict semantic validation shared by both engines."""

    root = _require_mapping(spec, "workbookSpec")
    _strict_keys(root, _SPEC_KEYS, "workbookSpec")
    if root.get("version") != WORKBOOK_SPEC_VERSION:
        raise _protocol("workbookSpec.version must be 1.", field_path="workbookSpec.version")
    policy = root.get("preservationPolicy")
    if action == "modify" and policy is None:
        raise _protocol(
            "Modify requires an explicit preservationPolicy.",
            field_path="workbookSpec.preservationPolicy",
        )
    policy = policy or "fail_on_unsupported_loss"
    if policy not in PRESERVATION_POLICIES:
        raise _protocol("The preservation policy is invalid.", field_path="workbookSpec.preservationPolicy")

    formats: dict[str, dict[str, Any]] = {}
    for index, raw_format in enumerate(_require_list(root.get("formats", []), "workbookSpec.formats")):
        item = _require_mapping(raw_format, f"workbookSpec.formats[{index}]")
        _strict_keys(item, _FORMAT_KEYS, f"workbookSpec.formats[{index}]")
        format_id = item.get("id")
        if not isinstance(format_id, str) or not format_id or format_id in formats:
            raise _protocol("Format IDs must be non-empty and unique.", field_path=f"workbookSpec.formats[{index}].id")
        formats[format_id] = dict(item)

    sheets = _require_list(root.get("sheets"), "workbookSpec.sheets")
    if not sheets or len(sheets) > 255:
        raise _protocol("A workbook requires between 1 and 255 sheets.", field_path="workbookSpec.sheets")
    names: list[str] = []
    folded_names: set[str] = set()
    visible = 0
    table_names: set[str] = set()
    for sheet_index, raw_sheet in enumerate(sheets):
        path = f"workbookSpec.sheets[{sheet_index}]"
        sheet = _require_mapping(raw_sheet, path)
        _strict_keys(sheet, _SHEET_KEYS, path)
        name = sheet.get("name")
        if (
            not isinstance(name, str)
            or not name
            or len(name) > 31
            or any(char in _SHEET_FORBIDDEN for char in name)
            or name.startswith("'")
            or name.endswith("'")
        ):
            raise _protocol("The sheet name is invalid.", field_path=f"{path}.name")
        if name.casefold() in folded_names:
            raise _protocol("Sheet names must be unique ignoring case.", field_path=f"{path}.name")
        folded_names.add(name.casefold())
        names.append(name)
        state = sheet.get("state")
        if state not in SHEET_STATES:
            raise _protocol("Every sheet requires state: visible, hidden, or veryHidden.", field_path=f"{path}.state")
        visible += int(state == "visible")
        _validate_sheet(sheet, path, formats, table_names)
    if action == "create" and visible == 0:
        raise _protocol("At least one workbook sheet must be visible.", field_path="workbookSpec.sheets")

    seen_ranges: set[str] = set()
    for index, raw_name in enumerate(_require_list(root.get("namedRanges", []), "workbookSpec.namedRanges")):
        path = f"workbookSpec.namedRanges[{index}]"
        item = _require_mapping(raw_name, path)
        _strict_keys(item, frozenset({"name", "formula"}), path)
        name = item.get("name")
        if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z_\\][A-Za-z0-9_.\\]{0,254}", name):
            raise _protocol("The named range name is invalid.", field_path=f"{path}.name")
        if name.casefold() in seen_ranges:
            raise _protocol("Named ranges must be unique ignoring case.", field_path=f"{path}.name")
        seen_ranges.add(name.casefold())
        _normal_formula(item.get("formula"), f"{path}.formula")

    properties = _require_mapping(root.get("properties", {}), "workbookSpec.properties")
    _strict_keys(
        properties,
        frozenset({"title", "subject", "author", "company", "category", "keywords", "comments"}),
        "workbookSpec.properties",
    )
    return WorkbookPlan(copy.deepcopy(dict(root)), policy, formats, tuple(names))


def _validate_sheet(
    sheet: Mapping[str, Any],
    path: str,
    formats: Mapping[str, Mapping[str, Any]],
    table_names: set[str],
) -> None:
    for row_index, row in enumerate(_require_list(sheet.get("data", []), f"{path}.data")):
        cells = _require_list(row, f"{path}.data[{row_index}]")
        if row_index >= 1_048_576 or len(cells) > 16_384:
            raise _protocol("Sheet data exceeds Excel bounds.", field_path=f"{path}.data[{row_index}]")
        for column_index, raw_cell in enumerate(cells):
            if isinstance(raw_cell, Mapping):
                cell_path = f"{path}.data[{row_index}][{column_index}]"
                _strict_keys(raw_cell, frozenset({"value", "formula", "formatId"}), cell_path)
                _validate_value_formula(raw_cell, cell_path)
                _validate_format_id(raw_cell.get("formatId"), formats, f"{cell_path}.formatId")
            elif not isinstance(raw_cell, (str, int, float, bool, type(None))):
                raise _protocol("Cell data must be a JSON scalar or cell object.", field_path=f"{path}.data")
    seen_cells: set[str] = set()
    for index, raw_cell in enumerate(_require_list(sheet.get("cells", []), f"{path}.cells")):
        cell_path = f"{path}.cells[{index}]"
        cell = _require_mapping(raw_cell, cell_path)
        _strict_keys(cell, _CELL_KEYS, cell_path)
        address = _validate_address(cell.get("address"), f"{cell_path}.address").replace("$", "")
        if address in seen_cells:
            raise _protocol("Cell addresses must be unique within explicit cells.", field_path=f"{cell_path}.address")
        seen_cells.add(address)
        _validate_value_formula(cell, cell_path)
        _validate_format_id(cell.get("formatId"), formats, f"{cell_path}.formatId")

    for index, raw_column in enumerate(_require_list(sheet.get("columns", []), f"{path}.columns")):
        item_path = f"{path}.columns[{index}]"
        column = _require_mapping(raw_column, item_path)
        _strict_keys(column, frozenset({"range", "width", "hidden", "formatId"}), item_path)
        if not isinstance(column.get("range"), str) or not _COLUMN_RANGE.fullmatch(column["range"]):
            raise _protocol("Column range is invalid.", field_path=f"{item_path}.range")
        _validate_format_id(column.get("formatId"), formats, f"{item_path}.formatId")
    for index, raw_row in enumerate(_require_list(sheet.get("rows", []), f"{path}.rows")):
        item_path = f"{path}.rows[{index}]"
        row = _require_mapping(raw_row, item_path)
        _strict_keys(row, frozenset({"index", "height", "hidden", "formatId"}), item_path)
        if not isinstance(row.get("index"), int) or not 1 <= row["index"] <= 1_048_576:
            raise _protocol("Row index is outside Excel bounds.", field_path=f"{item_path}.index")
        _validate_format_id(row.get("formatId"), formats, f"{item_path}.formatId")
    for index, merge in enumerate(_require_list(sheet.get("merges", []), f"{path}.merges")):
        _validate_range(merge, f"{path}.merges[{index}]")

    freeze = _require_mapping(sheet.get("freezePane", {}), f"{path}.freezePane")
    _strict_keys(freeze, frozenset({"row", "column"}), f"{path}.freezePane")
    if freeze:
        row = freeze.get("row", 0)
        column = freeze.get("column", 0)
        if not isinstance(row, int) or not 0 <= row <= 1_048_575:
            raise _protocol("freezePane.row is outside Excel bounds.", field_path=f"{path}.freezePane.row")
        if not isinstance(column, int) or not 0 <= column <= 16_383:
            raise _protocol("freezePane.column is outside Excel bounds.", field_path=f"{path}.freezePane.column")
    auto_filter = _require_mapping(sheet.get("autofilter", {}), f"{path}.autofilter")
    _strict_keys(auto_filter, frozenset({"range"}), f"{path}.autofilter")
    if auto_filter:
        _validate_range(auto_filter.get("range"), f"{path}.autofilter.range")

    for index, raw_table in enumerate(_require_list(sheet.get("tables", []), f"{path}.tables")):
        item_path = f"{path}.tables[{index}]"
        table = _require_mapping(raw_table, item_path)
        _strict_keys(
            table,
            frozenset({"name", "range", "style", "showHeaderRow", "showTotalRow", "bandedRows"}),
            item_path,
        )
        name = table.get("name")
        if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]{0,254}", name):
            raise _protocol("Table name is invalid.", field_path=f"{item_path}.name")
        if name.casefold() in table_names:
            raise _protocol("Table names must be workbook-wide unique.", field_path=f"{item_path}.name")
        table_names.add(name.casefold())
        _validate_range(table.get("range"), f"{item_path}.range")

    for index, raw_chart in enumerate(_require_list(sheet.get("charts", []), f"{path}.charts")):
        _validate_chart(_require_mapping(raw_chart, f"{path}.charts[{index}]"), f"{path}.charts[{index}]")
    for index, raw_validation in enumerate(_require_list(sheet.get("dataValidations", []), f"{path}.dataValidations")):
        _validate_data_validation(
            _require_mapping(raw_validation, f"{path}.dataValidations[{index}]"),
            f"{path}.dataValidations[{index}]",
        )
    for index, raw_conditional in enumerate(
        _require_list(sheet.get("conditionalFormats", []), f"{path}.conditionalFormats")
    ):
        _validate_conditional_format(
            _require_mapping(raw_conditional, f"{path}.conditionalFormats[{index}]"),
            f"{path}.conditionalFormats[{index}]",
            formats,
        )
    for index, raw_image in enumerate(_require_list(sheet.get("images", []), f"{path}.images")):
        item_path = f"{path}.images[{index}]"
        image = _require_mapping(raw_image, item_path)
        _strict_keys(
            image,
            frozenset({"artifactId", "position", "altText", "width", "height", "xOffset", "yOffset"}),
            item_path,
        )
        if not isinstance(image.get("artifactId"), str) or not image["artifactId"]:
            raise _protocol("Image artifactId is required.", field_path=f"{item_path}.artifactId")
        _validate_address(image.get("position"), f"{item_path}.position")
        if not isinstance(image.get("altText"), str) or not image["altText"]:
            raise _protocol("Image altText is required.", field_path=f"{item_path}.altText")
    _validate_page_setup(sheet.get("pageSetup", {}), f"{path}.pageSetup")


def _validate_value_formula(item: Mapping[str, Any], path: str) -> None:
    if "value" in item and "formula" in item:
        raise _protocol("A cell cannot define both value and formula.", field_path=path)
    if "formula" in item:
        _normal_formula(item["formula"], f"{path}.formula")
    if "value" in item and not isinstance(item["value"], (str, int, float, bool, type(None))):
        raise _protocol("Cell value must be a JSON scalar.", field_path=f"{path}.value")


def _validate_format_id(value: Any, formats: Mapping[str, Mapping[str, Any]], path: str) -> None:
    if value is not None and value not in formats:
        raise _protocol(f"Unknown formatId '{value}'.", field_path=path)


def _validate_chart(chart: Mapping[str, Any], path: str) -> None:
    _strict_keys(chart, frozenset({"type", "subtype", "title", "series", "position", "width", "height", "legend"}), path)
    if chart.get("type") not in _XLSXWRITER_CHART_TYPES:
        raise _unsupported("The requested chart type is not supported in Workbook Spec v1.", field_path=f"{path}.type")
    _validate_address(chart.get("position"), f"{path}.position")
    series = _require_list(chart.get("series"), f"{path}.series")
    if not series:
        raise _protocol("A chart requires at least one series.", field_path=f"{path}.series")
    for index, raw_series in enumerate(series):
        series_path = f"{path}.series[{index}]"
        item = _require_mapping(raw_series, series_path)
        _strict_keys(item, frozenset({"name", "categories", "values", "color"}), series_path)
        for key in ("categories", "values"):
            if key == "categories" and key not in item:
                continue
            reference = _require_mapping(item.get(key), f"{series_path}.{key}")
            _strict_keys(reference, frozenset({"sheet", "range"}), f"{series_path}.{key}")
            if not isinstance(reference.get("sheet"), str) or not reference["sheet"]:
                raise _protocol("Chart reference sheet is required.", field_path=f"{series_path}.{key}.sheet")
            _validate_range(reference.get("range"), f"{series_path}.{key}.range")


def _validate_data_validation(item: Mapping[str, Any], path: str) -> None:
    _strict_keys(
        item,
        frozenset(
            {
                "range", "type", "operator", "formula1", "formula2", "values", "allowBlank",
                "inputTitle", "inputMessage", "errorTitle", "errorMessage",
            }
        ),
        path,
    )
    _validate_range(item.get("range"), f"{path}.range")
    if item.get("type") not in {"any", "integer", "decimal", "list", "date", "time", "textLength", "custom"}:
        raise _unsupported("The requested data validation type is unsupported.", field_path=f"{path}.type")
    if item.get("type") == "list" and "values" not in item and "formula1" not in item:
        raise _protocol("List validation requires values or formula1.", field_path=path)
    for key in ("formula1", "formula2"):
        if isinstance(item.get(key), str):
            _normal_formula(item[key], f"{path}.{key}")


def _validate_conditional_format(
    item: Mapping[str, Any], path: str, formats: Mapping[str, Mapping[str, Any]]
) -> None:
    _strict_keys(item, frozenset({"range", "type", "operator", "value", "formula", "formatId"}), path)
    _validate_range(item.get("range"), f"{path}.range")
    if item.get("type") not in {
        "cell", "formula", "duplicate", "unique", "top", "bottom", "blanks", "noBlanks", "errors", "noErrors"
    }:
        raise _unsupported("The requested conditional-format type is unsupported.", field_path=f"{path}.type")
    _validate_format_id(item.get("formatId"), formats, f"{path}.formatId")
    if "formula" in item:
        _normal_formula(item["formula"], f"{path}.formula")


def _validate_page_setup(value: Any, path: str) -> None:
    page = _require_mapping(value, path)
    _strict_keys(
        page,
        frozenset({"orientation", "printArea", "margins", "header", "footer", "fitToWidth", "fitToHeight"}),
        path,
    )
    if "printArea" in page:
        _validate_range(page["printArea"], f"{path}.printArea")
    for key in ("header", "footer"):
        if key in page and (not isinstance(page[key], str) or len(page[key]) > 255):
            raise _protocol(f"{key} must be a string of at most 255 characters.", field_path=f"{path}.{key}")


@dataclass(slots=True)
class WorkbookExecutionResult:
    engine: str
    sheets: list[dict[str, Any]]
    tables: int
    charts: int
    warnings: list[dict[str, Any]] = field(default_factory=list)

    def public_dict(self) -> dict[str, Any]:
        return {
            "engine": self.engine,
            "sheets": self.sheets,
            "tables": self.tables,
            "charts": self.charts,
            "warnings": self.warnings,
        }


ImageResolver = Callable[[str], Path]


def create_workbook(
    spec: Mapping[str, Any], destination: Path | str, *, image_resolver: ImageResolver | None = None
) -> WorkbookExecutionResult:
    plan = validate_workbook_spec(spec, action="create")
    try:
        import xlsxwriter
    except ImportError:
        raise ArtifactError(
            "dependency_unavailable",
            "XlsxWriter is unavailable on the worker host.",
            "Install the pinned worker dependencies and retry.",
        ) from None

    workbook = xlsxwriter.Workbook(str(destination), {"constant_memory": False})
    try:
        workbook.set_properties(dict(plan.spec.get("properties", {})))
        formats = {format_id: workbook.add_format(_xlsxwriter_format(value)) for format_id, value in plan.formats.items()}
        for named in plan.spec.get("namedRanges", []):
            workbook.define_name(named["name"], _normal_formula(named["formula"], "namedRanges.formula"))
        sheet_summaries: list[dict[str, Any]] = []
        table_count = 0
        chart_count = 0
        worksheets = {
            sheet_spec["name"]: workbook.add_worksheet(sheet_spec["name"])
            for sheet_spec in plan.spec["sheets"]
        }
        visible_name = next(
            sheet_spec["name"]
            for sheet_spec in plan.spec["sheets"]
            if sheet_spec["state"] == "visible"
        )
        worksheets[visible_name].activate()
        for sheet_spec in plan.spec["sheets"]:
            worksheet = worksheets[sheet_spec["name"]]
            _apply_xlsxwriter_sheet(
                workbook,
                worksheet,
                sheet_spec,
                formats,
                plan.formats,
                image_resolver,
            )
            state = sheet_spec["state"]
            if state == "hidden":
                worksheet.hide()
            elif state == "veryHidden":
                worksheet.very_hidden()
            table_count += len(sheet_spec.get("tables", []))
            chart_count += len(sheet_spec.get("charts", []))
            sheet_summaries.append({"name": sheet_spec["name"], "state": state})
        close_result = workbook.close()
        if close_result:
            raise ArtifactError(
                "generation_failed",
                "XlsxWriter reported an output close failure.",
                "Check the output size and worker filesystem, then retry.",
            )
    except ArtifactError:
        try:
            workbook.close()
        except Exception:
            pass
        raise
    except Exception as exc:
        try:
            workbook.close()
        except Exception:
            pass
        raise ArtifactError(
            "generation_failed",
            "The workbook could not be generated from Workbook Spec v1.",
            "Correct the referenced range or supported feature and retry.",
            {"reason": type(exc).__name__},
        ) from None
    return WorkbookExecutionResult("xlsxwriter", sheet_summaries, table_count, chart_count)


def _xlsxwriter_format(spec: Mapping[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    font = spec.get("font", {})
    mapping = {
        "name": "font_name", "size": "font_size", "bold": "bold", "italic": "italic",
        "underline": "underline", "color": "font_color",
    }
    for source, target in mapping.items():
        if source in font:
            result[target] = font[source]
    fill = spec.get("fill", {})
    if fill.get("pattern", "solid") != "none" and "color" in fill:
        result["bg_color"] = fill["color"]
        result["pattern"] = 1
    alignment = spec.get("alignment", {})
    for source, target in {
        "horizontal": "align", "vertical": "valign", "wrapText": "text_wrap", "textRotation": "rotation"
    }.items():
        if source in alignment:
            result[target] = alignment[source]
    border = spec.get("border", {})
    border_styles = {"none": 0, "thin": 1, "medium": 2, "dashed": 3, "dotted": 4, "thick": 5, "double": 6}
    if "style" in border:
        result["border"] = border_styles[border["style"]]
    if "color" in border:
        result["border_color"] = border["color"]
    if "numberFormat" in spec:
        result["num_format"] = spec["numberFormat"]
    if "locked" in spec:
        result["locked"] = spec["locked"]
    return result


def _apply_xlsxwriter_sheet(
    workbook: Any,
    worksheet: Any,
    spec: Mapping[str, Any],
    formats: Mapping[str, Any],
    format_specs: Mapping[str, Mapping[str, Any]],
    image_resolver: ImageResolver | None,
) -> None:
    cell_formats: dict[tuple[str | None, str], Any] = {}
    if "tabColor" in spec:
        worksheet.set_tab_color(spec["tabColor"])
    for merge in spec.get("merges", []):
        worksheet.merge_range(merge, "")
    for row_index, row in enumerate(spec.get("data", [])):
        for column_index, item in enumerate(row):
            _xlsxwriter_write(workbook, worksheet, row_index, column_index, item, formats, format_specs, cell_formats)
    for cell in spec.get("cells", []):
        _xlsxwriter_write(workbook, worksheet, cell["address"], None, cell, formats, format_specs, cell_formats)
        if "comment" in cell:
            worksheet.write_comment(cell["address"], cell["comment"])
        if "hyperlink" in cell:
            display = cell.get("value") if isinstance(cell.get("value"), str) else cell["hyperlink"]
            worksheet.write_url(cell["address"], cell["hyperlink"], formats.get(cell.get("formatId")), display)
    for column in spec.get("columns", []):
        options = {"hidden": column.get("hidden", False)}
        column_range = column["range"]
        if ":" not in column_range:
            column_range = f"{column_range}:{column_range}"
        worksheet.set_column(column_range, column.get("width"), formats.get(column.get("formatId")), options)
    for row in spec.get("rows", []):
        options = {"hidden": row.get("hidden", False)}
        worksheet.set_row(row["index"] - 1, row.get("height"), formats.get(row.get("formatId")), options)
    freeze = spec.get("freezePane")
    if freeze:
        worksheet.freeze_panes(freeze.get("row", 0), freeze.get("column", 0))
    auto_filter = spec.get("autofilter")
    table_ranges = {table["range"] for table in spec.get("tables", [])}
    if auto_filter and auto_filter["range"] not in table_ranges:
        worksheet.autofilter(auto_filter["range"])
    for table in spec.get("tables", []):
        options = {
            "name": table["name"],
            "style": table.get("style", "Table Style Medium 2"),
            "header_row": table.get("showHeaderRow", True),
            "total_row": table.get("showTotalRow", False),
            "banded_rows": table.get("bandedRows", True),
        }
        headers = _xlsxwriter_table_headers(spec, table["range"])
        if headers is not None:
            # XlsxWriter otherwise overwrites declared header cells with its
            # generated Column1/Column2 labels.
            options["columns"] = [{"header": header} for header in headers]
        worksheet.add_table(table["range"], options)
    for chart_spec in spec.get("charts", []):
        options = {"type": chart_spec["type"]}
        if chart_spec.get("subtype"):
            options["subtype"] = chart_spec["subtype"]
        chart = workbook.add_chart(options)
        for series in chart_spec["series"]:
            series_options: dict[str, Any] = {"values": _xlsxwriter_ref(series["values"])}
            if "categories" in series:
                series_options["categories"] = _xlsxwriter_ref(series["categories"])
            if "name" in series:
                series_options["name"] = series["name"]
            if "color" in series:
                series_options["fill"] = {"color": series["color"]}
                series_options["line"] = {"color": series["color"]}
            chart.add_series(series_options)
        if "title" in chart_spec:
            chart.set_title({"name": chart_spec["title"]})
        if chart_spec.get("legend") == "none":
            chart.set_legend({"none": True})
        elif chart_spec.get("legend"):
            chart.set_legend({"position": chart_spec["legend"]})
        worksheet.insert_chart(
            chart_spec["position"],
            chart,
            {key: chart_spec[key] for key in ("width", "height") if key in chart_spec},
        )
    for validation in spec.get("dataValidations", []):
        worksheet.data_validation(validation["range"], _xlsxwriter_validation(validation))
    for conditional in spec.get("conditionalFormats", []):
        worksheet.conditional_format(conditional["range"], _xlsxwriter_conditional(conditional, formats))
    for image in spec.get("images", []):
        if image_resolver is None:
            raise _unsupported("Image insertion requires a trusted artifact resolver.", field_path="images.artifactId")
        image_path = Path(image_resolver(image["artifactId"]))
        options: dict[str, Any] = {
            "description": image["altText"],
            "decorative": False,
            "x_offset": image.get("xOffset", 0),
            "y_offset": image.get("yOffset", 0),
        }
        if "width" in image or "height" in image:
            try:
                from PIL import Image
                with Image.open(image_path) as picture:
                    width, height = picture.size
            except Exception:
                raise ArtifactError(
                    "input_format_invalid",
                    "An image artifact could not be decoded.",
                    "Attach a valid PNG or JPEG image.",
                ) from None
            if "width" in image:
                options["x_scale"] = image["width"] / width
            if "height" in image:
                options["y_scale"] = image["height"] / height
        worksheet.insert_image(image["position"], str(image_path), options)
    _apply_xlsxwriter_page(worksheet, spec.get("pageSetup", {}))


def _xlsxwriter_table_headers(spec: Mapping[str, Any], table_range: str) -> list[str] | None:
    """Project declared header values into XlsxWriter's native table options."""

    from xlsxwriter.utility import xl_cell_to_rowcol

    endpoints = table_range.split(":", 1)
    first = endpoints[0]
    last = endpoints[-1]
    first_row, first_column = xl_cell_to_rowcol(first)
    _, last_column = xl_cell_to_rowcol(last)
    declared: dict[tuple[int, int], Any] = {}
    for row_index, row in enumerate(spec.get("data", [])):
        for column_index, item in enumerate(row):
            declared[(row_index, column_index)] = item
    for item in spec.get("cells", []):
        row_index, column_index = xl_cell_to_rowcol(item["address"])
        declared[(row_index, column_index)] = item

    headers: list[str] = []
    any_declared = False
    for offset, column_index in enumerate(range(first_column, last_column + 1), start=1):
        coordinate = (first_row, column_index)
        item = declared.get(coordinate)
        if isinstance(item, Mapping):
            value = item.get("formula") if "formula" in item else item.get("value")
            any_declared = any_declared or "formula" in item or "value" in item
        else:
            value = item
            any_declared = any_declared or coordinate in declared
        headers.append(str(value) if value not in {None, ""} else f"Column{offset}")
    return headers if any_declared else None


def _xlsxwriter_write(
    workbook: Any,
    worksheet: Any,
    row_or_address: int | str,
    column: int | None,
    item: Any,
    formats: Mapping[str, Any],
    format_specs: Mapping[str, Mapping[str, Any]],
    cell_formats: dict[tuple[str | None, str], Any],
) -> None:
    if isinstance(item, Mapping):
        format_id = item.get("formatId")
        format_object = formats.get(format_id)
        if item.get("numberFormat"):
            cache_key = (format_id, item["numberFormat"])
            if cache_key not in cell_formats:
                properties = _xlsxwriter_format(format_specs.get(format_id, {}))
                properties["num_format"] = item["numberFormat"]
                cell_formats[cache_key] = workbook.add_format(properties)
            format_object = cell_formats[cache_key]
        if "formula" in item:
            value = _normal_formula(item["formula"], "formula")
            if column is None:
                worksheet.write_formula(row_or_address, value, format_object)
            else:
                worksheet.write_formula(row_or_address, column, value, format_object)
        else:
            value = item.get("value")
            if column is None:
                if isinstance(value, str) and value.startswith("="):
                    worksheet.write_string(row_or_address, value, format_object)
                else:
                    worksheet.write(row_or_address, value, format_object)
            else:
                if isinstance(value, str) and value.startswith("="):
                    worksheet.write_string(row_or_address, column, value, format_object)
                else:
                    worksheet.write(row_or_address, column, value, format_object)
        return
    if column is None:
        if isinstance(item, str) and item.startswith("="):
            worksheet.write_string(row_or_address, item)
        else:
            worksheet.write(row_or_address, item)
    else:
        if isinstance(item, str) and item.startswith("="):
            worksheet.write_string(row_or_address, column, item)
        else:
            worksheet.write(row_or_address, column, item)


def _xlsxwriter_ref(reference: Mapping[str, Any]) -> str:
    escaped = reference["sheet"].replace("'", "''")
    return f"='{escaped}'!{reference['range']}"


def _xlsxwriter_validation(spec: Mapping[str, Any]) -> dict[str, Any]:
    type_map = {"integer": "integer", "decimal": "decimal", "textLength": "length", "custom": "custom"}
    result: dict[str, Any] = {"validate": type_map.get(spec["type"], spec["type"])}
    operator_map = {
        "notBetween": "not between", "notEqual": "!=", "greaterThan": ">", "lessThan": "<",
        "greaterThanOrEqual": ">=", "lessThanOrEqual": "<=", "equal": "==",
    }
    if "operator" in spec:
        result["criteria"] = operator_map.get(spec["operator"], spec["operator"])
    if "values" in spec:
        result["source"] = spec["values"]
    elif "formula1" in spec:
        result["value"] = spec["formula1"]
    if "formula2" in spec:
        result["maximum"] = spec["formula2"]
    for source, target in {
        "allowBlank": "ignore_blank", "inputTitle": "input_title", "inputMessage": "input_message",
        "errorTitle": "error_title", "errorMessage": "error_message",
    }.items():
        if source in spec:
            result[target] = spec[source]
    return result


def _xlsxwriter_conditional(spec: Mapping[str, Any], formats: Mapping[str, Any]) -> dict[str, Any]:
    type_map = {
        "cell": "cell", "formula": "formula", "duplicate": "duplicate", "unique": "unique",
        "top": "top", "bottom": "bottom", "blanks": "blanks", "noBlanks": "no_blanks",
        "errors": "errors", "noErrors": "no_errors",
    }
    result: dict[str, Any] = {"type": type_map[spec["type"]], "format": formats[spec["formatId"]]}
    if "operator" in spec:
        result["criteria"] = {
            "notBetween": "not between", "notEqual": "!=", "greaterThan": ">", "lessThan": "<",
            "greaterThanOrEqual": ">=", "lessThanOrEqual": "<=", "equal": "==",
        }.get(spec["operator"], spec["operator"])
    if "value" in spec:
        result["value"] = spec["value"]
    if "formula" in spec:
        result["criteria"] = _normal_formula(spec["formula"], "conditionalFormats.formula")
    return result


def _apply_xlsxwriter_page(worksheet: Any, page: Mapping[str, Any]) -> None:
    if page.get("orientation") == "landscape":
        worksheet.set_landscape()
    elif page.get("orientation") == "portrait":
        worksheet.set_portrait()
    if "printArea" in page:
        worksheet.print_area(page["printArea"])
    if "fitToWidth" in page or "fitToHeight" in page:
        worksheet.fit_to_pages(page.get("fitToWidth", 1), page.get("fitToHeight", 0))
    margins = page.get("margins", {})
    if margins:
        worksheet.set_margins(margins.get("left", 0.7), margins.get("right", 0.7), margins.get("top", 0.75), margins.get("bottom", 0.75))
        if "header" in margins or "footer" in margins:
            worksheet.set_header("", {"margin": margins.get("header", 0.3)})
            worksheet.set_footer("", {"margin": margins.get("footer", 0.3)})
    if "header" in page:
        worksheet.set_header(page["header"])
    if "footer" in page:
        worksheet.set_footer(page["footer"])


def modify_workbook(
    spec: Mapping[str, Any], staged_workbook: Path | str, *, image_resolver: ImageResolver | None = None
) -> WorkbookExecutionResult:
    """Apply declarative edits to a private source copy using openpyxl."""

    plan = validate_workbook_spec(spec, action="modify")
    try:
        import openpyxl
        from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, PieChart, Reference, ScatterChart
        from openpyxl.chart.series_factory import SeriesFactory
        from openpyxl.comments import Comment
        from openpyxl.formatting.rule import FormulaRule, Rule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        raise ArtifactError(
            "dependency_unavailable",
            "openpyxl is unavailable on the worker host.",
            "Install the pinned worker dependencies and retry.",
        ) from None

    path = Path(staged_workbook)
    try:
        workbook = openpyxl.load_workbook(path, keep_vba=False, keep_links=True)
    except Exception:
        raise ArtifactError(
            "input_format_invalid",
            "The source workbook could not be opened as editable XLSX.",
            "Remove encryption or corruption and attach a standard .xlsx workbook.",
        ) from None

    styles = {format_id: _openpyxl_style(value, Font, PatternFill, Alignment, Border, Side, Protection) for format_id, value in plan.formats.items()}
    tables_added = 0
    charts_added = 0
    try:
        for key, value in plan.spec.get("properties", {}).items():
            target = "description" if key == "comments" else key
            if hasattr(workbook.properties, target):
                setattr(workbook.properties, target, value)
        existing_names = {name.casefold() for name in workbook.sheetnames}
        for sheet_spec in plan.spec["sheets"]:
            name = sheet_spec["name"]
            if name in workbook.sheetnames:
                worksheet = workbook[name]
            elif name.casefold() in existing_names:
                raise _protocol("Sheet names conflict ignoring case.", field_path=f"sheets.{name}")
            else:
                worksheet = workbook.create_sheet(name)
                existing_names.add(name.casefold())
            worksheet.sheet_state = sheet_spec["state"]
            if "tabColor" in sheet_spec:
                worksheet.sheet_properties.tabColor = sheet_spec["tabColor"].lstrip("#")
            _apply_openpyxl_cells(worksheet, sheet_spec, styles, Comment)
            _apply_openpyxl_dimensions(worksheet, sheet_spec, styles)
            for merge in sheet_spec.get("merges", []):
                if str(merge) not in {str(value) for value in worksheet.merged_cells.ranges}:
                    worksheet.merge_cells(merge)
            freeze = sheet_spec.get("freezePane")
            if freeze:
                worksheet.freeze_panes = worksheet.cell(freeze.get("row", 0) + 1, freeze.get("column", 0) + 1)
            if sheet_spec.get("autofilter"):
                worksheet.auto_filter.ref = sheet_spec["autofilter"]["range"]
            for table_spec in sheet_spec.get("tables", []):
                table = Table(displayName=table_spec["name"], ref=table_spec["range"])
                table.tableStyleInfo = TableStyleInfo(
                    name=table_spec.get("style", "TableStyleMedium2").replace(" ", ""),
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=table_spec.get("bandedRows", True),
                    showColumnStripes=False,
                )
                table.headerRowCount = 1 if table_spec.get("showHeaderRow", True) else 0
                table.totalsRowShown = table_spec.get("showTotalRow", False)
                worksheet.add_table(table)
                tables_added += 1
            for chart_spec in sheet_spec.get("charts", []):
                chart = _openpyxl_chart(
                    workbook, chart_spec, AreaChart, BarChart, DoughnutChart, LineChart, PieChart,
                    ScatterChart, Reference, SeriesFactory,
                )
                worksheet.add_chart(chart, chart_spec["position"])
                charts_added += 1
            for validation_spec in sheet_spec.get("dataValidations", []):
                validation = _openpyxl_validation(validation_spec, DataValidation)
                worksheet.add_data_validation(validation)
                validation.add(validation_spec["range"])
            for conditional in sheet_spec.get("conditionalFormats", []):
                rule = _openpyxl_conditional(conditional, styles, FormulaRule, Rule)
                worksheet.conditional_formatting.add(conditional["range"], rule)
            for image in sheet_spec.get("images", []):
                if image_resolver is None:
                    raise _unsupported("Image insertion requires a trusted artifact resolver.", field_path="images.artifactId")
                from openpyxl.drawing.image import Image as OpenpyxlImage
                picture = OpenpyxlImage(str(image_resolver(image["artifactId"])))
                if "width" in image:
                    picture.width = image["width"]
                if "height" in image:
                    picture.height = image["height"]
                worksheet.add_image(picture, image["position"])
            _apply_openpyxl_page(worksheet, sheet_spec.get("pageSetup", {}))
        for named in plan.spec.get("namedRanges", []):
            from openpyxl.workbook.defined_name import DefinedName
            workbook.defined_names[named["name"]] = DefinedName(named["name"], attr_text=named["formula"].lstrip("="))
        if not any(sheet.sheet_state == "visible" for sheet in workbook.worksheets):
            raise _protocol("At least one workbook sheet must remain visible.", field_path="workbookSpec.sheets")
        workbook.save(path)
    except ArtifactError:
        workbook.close()
        raise
    except Exception as exc:
        workbook.close()
        raise ArtifactError(
            "modification_failed",
            "The source workbook could not be modified from Workbook Spec v1.",
            "Correct the supported edit or conflicting workbook feature and retry.",
            {"reason": type(exc).__name__},
        ) from None
    workbook.close()
    return WorkbookExecutionResult(
        "openpyxl",
        [{"name": item["name"], "state": item["state"]} for item in plan.spec["sheets"]],
        tables_added,
        charts_added,
    )


def _apply_openpyxl_cells(worksheet: Any, spec: Mapping[str, Any], styles: Mapping[str, Mapping[str, Any]], Comment: Any) -> None:
    for row_index, row in enumerate(spec.get("data", []), start=1):
        for column_index, item in enumerate(row, start=1):
            cell = worksheet.cell(row_index, column_index)
            _openpyxl_set_cell(cell, item, styles)
    for item in spec.get("cells", []):
        cell = worksheet[item["address"]]
        _openpyxl_set_cell(cell, item, styles)
        if "numberFormat" in item:
            cell.number_format = item["numberFormat"]
        if "comment" in item:
            cell.comment = Comment(item["comment"], "Orca")
        if "hyperlink" in item:
            cell.hyperlink = item["hyperlink"]


def _openpyxl_set_cell(cell: Any, item: Any, styles: Mapping[str, Mapping[str, Any]]) -> None:
    if isinstance(item, Mapping):
        cell.value = _normal_formula(item["formula"], "formula") if "formula" in item else item.get("value")
        if item.get("formatId"):
            for attribute, value in styles[item["formatId"]].items():
                setattr(cell, attribute, copy.copy(value))
    else:
        cell.value = item
    if isinstance(cell.value, str) and cell.value.startswith("=") and not (
        isinstance(item, Mapping) and "formula" in item
    ):
        cell.data_type = "s"


def _openpyxl_style(spec: Mapping[str, Any], Font: Any, PatternFill: Any, Alignment: Any, Border: Any, Side: Any, Protection: Any) -> dict[str, Any]:
    result: dict[str, Any] = {}
    font = spec.get("font", {})
    if font:
        result["font"] = Font(
            name=font.get("name"), size=font.get("size"), bold=font.get("bold", False),
            italic=font.get("italic", False), underline="single" if font.get("underline") else None,
            color=str(font.get("color", "")).lstrip("#") or None,
        )
    fill = spec.get("fill", {})
    if fill:
        result["fill"] = PatternFill(fill_type=fill.get("pattern", "solid"), fgColor=fill.get("color", "").lstrip("#"))
    alignment = spec.get("alignment", {})
    if alignment:
        result["alignment"] = Alignment(
            horizontal=alignment.get("horizontal"), vertical=alignment.get("vertical"),
            wrap_text=alignment.get("wrapText"), text_rotation=alignment.get("textRotation", 0),
        )
    border = spec.get("border", {})
    if border:
        side = Side(style=None if border.get("style") == "none" else border.get("style"), color=border.get("color", "").lstrip("#") or None)
        result["border"] = Border(left=side, right=side, top=side, bottom=side)
    if "numberFormat" in spec:
        result["number_format"] = spec["numberFormat"]
    if "locked" in spec:
        result["protection"] = Protection(locked=spec["locked"])
    return result


def _apply_openpyxl_dimensions(worksheet: Any, spec: Mapping[str, Any], styles: Mapping[str, Mapping[str, Any]]) -> None:
    from openpyxl.utils.cell import column_index_from_string, get_column_letter
    for column in spec.get("columns", []):
        start, end = (column["range"].split(":") * 2)[:2]
        for index in range(column_index_from_string(start), column_index_from_string(end) + 1):
            dimension = worksheet.column_dimensions[get_column_letter(index)]
            if "width" in column:
                dimension.width = column["width"]
            dimension.hidden = column.get("hidden", False)
            if column.get("formatId"):
                for attribute, value in styles[column["formatId"]].items():
                    setattr(dimension, attribute, copy.copy(value))
    for row in spec.get("rows", []):
        dimension = worksheet.row_dimensions[row["index"]]
        if "height" in row:
            dimension.height = row["height"]
        dimension.hidden = row.get("hidden", False)
        if row.get("formatId"):
            for attribute, value in styles[row["formatId"]].items():
                setattr(dimension, attribute, copy.copy(value))


def _openpyxl_chart(workbook: Any, spec: Mapping[str, Any], AreaChart: Any, BarChart: Any, DoughnutChart: Any, LineChart: Any, PieChart: Any, ScatterChart: Any, Reference: Any, SeriesFactory: Any) -> Any:
    classes = {
        "area": AreaChart, "bar": BarChart, "column": BarChart, "doughnut": DoughnutChart,
        "line": LineChart, "pie": PieChart, "scatter": ScatterChart,
    }
    chart = classes[spec["type"]]()
    if spec["type"] == "column":
        chart.type = "col"
    elif spec["type"] == "bar":
        chart.type = "bar"
    if "title" in spec:
        chart.title = spec["title"]
    chart.width = spec.get("width", 576) / 96
    chart.height = spec.get("height", 336) / 96
    if spec.get("legend") == "none":
        chart.legend = None
    elif spec.get("legend"):
        chart.legend.position = {"top": "t", "bottom": "b", "left": "l", "right": "r"}[spec["legend"]]
    for series_spec in spec["series"]:
        values = _reference_from_spec(workbook, series_spec["values"], Reference)
        title = series_spec.get("name")
        series = SeriesFactory(values, title=title)
        if "categories" in series_spec:
            categories = _reference_from_spec(workbook, series_spec["categories"], Reference)
            series.cat = categories
        chart.series.append(series)
    return chart


def _reference_from_spec(workbook: Any, spec: Mapping[str, Any], Reference: Any) -> Any:
    from openpyxl.utils.cell import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(spec["range"])
    return Reference(workbook[spec["sheet"]], min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)


def _openpyxl_validation(spec: Mapping[str, Any], DataValidation: Any) -> Any:
    type_map = {"integer": "whole", "textLength": "textLength", "custom": "custom"}
    formula1: Any = spec.get("formula1")
    if "values" in spec:
        formula1 = '"' + ",".join(str(value) for value in spec["values"]) + '"'
    return DataValidation(
        type=type_map.get(spec["type"], spec["type"]), operator=spec.get("operator"),
        formula1=formula1, formula2=spec.get("formula2"), allow_blank=spec.get("allowBlank", False),
        promptTitle=spec.get("inputTitle"), prompt=spec.get("inputMessage"),
        errorTitle=spec.get("errorTitle"), error=spec.get("errorMessage"),
    )


def _openpyxl_conditional(spec: Mapping[str, Any], styles: Mapping[str, Mapping[str, Any]], FormulaRule: Any, Rule: Any) -> Any:
    style = styles.get(spec.get("formatId"), {})
    differential = {key: style[key] for key in ("font", "fill", "border") if key in style}
    if spec["type"] == "formula":
        return FormulaRule(formula=[spec.get("formula", "")], **differential)
    operator_map = {"notBetween": "notBetween", "notEqual": "notEqual", "greaterThan": "greaterThan", "lessThan": "lessThan", "greaterThanOrEqual": "greaterThanOrEqual", "lessThanOrEqual": "lessThanOrEqual", "equal": "equal"}
    rule_type = {"cell": "cellIs", "duplicate": "duplicateValues", "unique": "uniqueValues", "top": "top10", "bottom": "top10", "blanks": "containsBlanks", "noBlanks": "notContainsBlanks", "errors": "containsErrors", "noErrors": "notContainsErrors"}[spec["type"]]
    formula = [str(spec["value"])] if "value" in spec else []
    rule = Rule(type=rule_type, operator=operator_map.get(spec.get("operator")), formula=formula)
    if spec["type"] == "bottom":
        rule.bottom = True
    for key, value in differential.items():
        setattr(rule.dxf, key, value)
    return rule


def _apply_openpyxl_page(worksheet: Any, page: Mapping[str, Any]) -> None:
    if "orientation" in page:
        worksheet.page_setup.orientation = page["orientation"]
    if "printArea" in page:
        worksheet.print_area = page["printArea"]
    if "fitToWidth" in page or "fitToHeight" in page:
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = page.get("fitToWidth", 1)
        worksheet.page_setup.fitToHeight = page.get("fitToHeight", 0)
    margins = page.get("margins", {})
    for key in ("left", "right", "top", "bottom", "header", "footer"):
        if key in margins:
            setattr(worksheet.page_margins, key, margins[key])
    if "header" in page:
        worksheet.oddHeader.center.text = page["header"]
    if "footer" in page:
        worksheet.oddFooter.center.text = page["footer"]
