"""Streaming XLSX cell inspection/extraction for workbooks too large to DOM-parse."""

from __future__ import annotations

import datetime
import re
from pathlib import Path
from typing import Any, Mapping

import openpyxl
from defusedxml.ElementTree import fromstring as defused_fromstring

from orca_excel_artifact.archive import inspect_zip, safe_read_member
from orca_excel_artifact.artifacts import sha256_file
from orca_excel_artifact.errors import ArtifactError

MAX_SCANNED_CELLS = 8_000_000
MAX_EXTRACT_LIMIT = 200
MAX_CURSOR = 10_000_000
MAX_ITEM_TEXT = 32_767
_MACRO_MARKERS = re.compile(r"macroEnabled|vbaProject|activeX|oleObject", re.IGNORECASE)
_PERCENT_DECIMALS = re.compile(r"0\.(0+)%")


def _limit_error(message: str, recovery: str) -> ArtifactError:
    return ArtifactError("archive_limit_exceeded", message, recovery)


def _validated_source(value: Mapping[str, Any]) -> Path:
    path = Path(str(value.get("sourcePath")))
    if not path.is_absolute() or path.is_symlink():
        raise ArtifactError(
            "protocol_invalid",
            "The workbook source path must be an absolute regular path.",
            "Retry through a supported Orca build.",
        )
    digest, _size = sha256_file(path)
    expected = value.get("expectedSha256")
    if not isinstance(expected, str) or digest != expected.casefold():
        raise ArtifactError(
            "input_hash_mismatch",
            "The workbook input hash does not match the admitted attachment.",
            "Attach the current workbook and retry.",
        )
    return path


def _validate_package(path: Path) -> None:
    archive = inspect_zip(path)
    content_types = safe_read_member(archive, "[Content_Types].xml", max_bytes=1024 * 1024)
    if b"spreadsheetml.sheet.main+xml" not in content_types:
        raise ArtifactError(
            "input_format_invalid",
            "The file content is not an XLSX workbook.",
            "Export the data as a standard .xlsx workbook and retry.",
        )
    if _MACRO_MARKERS.search(content_types.decode("utf-8", "replace")):
        raise ArtifactError(
            "input_format_invalid",
            "Macro, ActiveX, and embedded OLE workbook content is not supported.",
            "Re-save the workbook as a standard non-macro .xlsx file.",
        )
    for name in archive.names:
        if not name.endswith(".rels"):
            continue
        payload = safe_read_member(archive, name, max_bytes=4 * 1024 * 1024)
        if b"<!DOCTYPE" in payload or b"<!ENTITY" in payload:
            raise ArtifactError(
                "input_format_invalid",
                "The workbook contains a forbidden XML declaration.",
                "Re-save the workbook with a trusted spreadsheet application.",
            )
        for node in defused_fromstring(payload.decode("utf-8")).iter():
            if node.get("TargetMode") == "External":
                raise ArtifactError(
                    "input_format_invalid",
                    "External XLSX relationships are not supported.",
                    "Remove external links and re-save the workbook.",
                )


def _number_format(cell: Any) -> dict[str, str]:
    fmt = cell.number_format
    return {"numberFormat": fmt} if isinstance(fmt, str) and fmt and fmt != "General" else {}


def _number_text(value: int | float, number_format: str | None) -> str:
    raw = str(value) if isinstance(value, int) else repr(value)
    if number_format and "%" in number_format:
        decimals_match = _PERCENT_DECIMALS.search(number_format)
        decimals = len(decimals_match.group(1)) if decimals_match else 0
        return f"{value * 100:.{decimals}f}%"
    return raw


def _formatted_value(value: Any, number_format: str | None) -> tuple[str, str]:
    """Return (text, rawValue) for a cached or plain scalar value."""

    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE", "1" if value else "0")
    if isinstance(value, str):
        return (value[:MAX_ITEM_TEXT], value[:MAX_ITEM_TEXT])
    if isinstance(value, datetime.datetime):
        midnight = (value.hour, value.minute, value.second, value.microsecond) == (0, 0, 0, 0)
        return (value.date().isoformat() if midnight else value.isoformat(), value.isoformat())
    if isinstance(value, (datetime.date, datetime.time)):
        return (value.isoformat(), value.isoformat())
    if isinstance(value, (int, float)):
        raw = str(value) if isinstance(value, int) else repr(value)
        return (_number_text(value, number_format), raw)
    text = str(value)[:MAX_ITEM_TEXT]
    return (text, text)


def _cell_item(sheet: str, cell: Any) -> dict[str, Any]:
    value = cell.value
    base: dict[str, Any] = {"kind": "xlsx_cell", "sheet": sheet, "cell": cell.coordinate}
    if cell.data_type == "f":
        source = value.text if hasattr(value, "text") else str(value)
        formula = source[1:] if source.startswith("=") else source
        return {
            **base,
            "text": f"={formula}" if formula else "=SHARED_FORMULA",
            "valueType": "formula",
            **({"formula": formula[:MAX_ITEM_TEXT]} if formula else {}),
            **_number_format(cell),
        }
    if isinstance(value, bool):
        return {**base, "text": "TRUE" if value else "FALSE", "valueType": "boolean", "rawValue": "1" if value else "0"}
    if cell.data_type == "e":
        return {**base, "text": str(value), "valueType": "error", "rawValue": str(value)}
    if isinstance(value, str):
        return {**base, "text": value[:MAX_ITEM_TEXT], "valueType": "text"}
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        text, raw = _formatted_value(value, None)
        return {**base, "text": text, "valueType": "date", "rawValue": raw, **_number_format(cell)}
    if isinstance(value, (int, float)):
        fmt = cell.number_format if isinstance(cell.number_format, str) else None
        text, raw = (_number_text(value, fmt), str(value) if isinstance(value, int) else repr(value))
        return {**base, "text": text, "valueType": "number", "rawValue": raw, **_number_format(cell)}
    text, raw = _formatted_value(value, None)
    return {**base, "text": text, "valueType": "text", "rawValue": raw}


def _open_streaming_workbook(path: Path, *, data_only: bool) -> tuple[Any, Any]:
    """Open by handle: artifact payload files have no .xlsx extension for openpyxl."""

    stream = path.open("rb")
    try:
        return openpyxl.load_workbook(stream, read_only=True, data_only=data_only), stream
    except Exception:
        stream.close()
        raise


def _fill_cached_formula_values(
    path: Path, items: list[dict[str, Any]], pending: list[tuple[int, str, int, int]]
) -> None:
    by_sheet: dict[str, dict[tuple[int, int], int]] = {}
    for item_index, sheet, row, column in pending:
        by_sheet.setdefault(sheet, {})[(row, column)] = item_index
    workbook, stream = _open_streaming_workbook(path, data_only=True)
    try:
        for worksheet in workbook.worksheets:
            wanted = by_sheet.get(worksheet.title)
            if not wanted:
                continue
            last_row = max(row for row, _column in wanted)
            for row in worksheet.iter_rows(max_row=last_row):
                for cell in row:
                    if cell.value is None:
                        continue
                    item_index = wanted.get((cell.row, cell.column))
                    if item_index is None:
                        continue
                    item = items[item_index]
                    text, raw = _formatted_value(cell.value, item.get("numberFormat"))
                    item["text"] = text
                    item["rawValue"] = raw
    finally:
        workbook.close()
        stream.close()


def _scan(
    path: Path, *, collect_window: bool, cursor: int, limit: int
) -> dict[str, Any]:
    workbook, stream = _open_streaming_workbook(path, data_only=False)
    sheets: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []
    pending_formulas: list[tuple[int, str, int, int]] = []
    next_cursor: int | None = None
    index = 0
    scanned = 0
    try:
        for worksheet in workbook.worksheets:
            counts = {
                "name": worksheet.title,
                "cellCount": 0,
                "textCellCount": 0,
                "numericCellCount": 0,
                "formulaCellCount": 0,
            }
            for row in worksheet.iter_rows():
                for cell in row:
                    scanned += 1
                    if scanned > MAX_SCANNED_CELLS:
                        raise _limit_error(
                            "The workbook exceeds the streaming cell scan limit.",
                            "Extract one smaller sheet or export the needed range separately.",
                        )
                    value = cell.value
                    if value is None or value == "":
                        continue
                    counts["cellCount"] += 1
                    if cell.data_type == "f":
                        counts["formulaCellCount"] += 1
                    elif isinstance(value, bool):
                        pass
                    elif isinstance(value, str) and cell.data_type != "e":
                        counts["textCellCount"] += 1
                    elif isinstance(value, (int, float)):
                        counts["numericCellCount"] += 1
                    if collect_window and next_cursor is None and index >= cursor:
                        if len(items) < limit:
                            items.append(_cell_item(worksheet.title, cell))
                            if cell.data_type == "f":
                                pending_formulas.append(
                                    (len(items) - 1, worksheet.title, cell.row, cell.column)
                                )
                        else:
                            next_cursor = index
                    index += 1
            sheets.append(counts)
    finally:
        workbook.close()
        stream.close()
    if pending_formulas:
        _fill_cached_formula_values(path, items, pending_formulas)
    result: dict[str, Any] = {"sheets": sheets}
    if collect_window:
        result["items"] = items
        if next_cursor is not None:
            result["nextCursor"] = next_cursor
    return result


def run_xlsx_extraction_job(value: Mapping[str, Any]) -> dict[str, Any]:
    action = str(value.get("action"))
    source = _validated_source(value)
    _validate_package(source)
    if action == "inspect_xlsx":
        return {"ok": True, "action": action, **_scan(source, collect_window=False, cursor=0, limit=0)}
    cursor = value.get("cursor", 0)
    limit = value.get("limit", MAX_EXTRACT_LIMIT)
    if (
        not isinstance(cursor, int)
        or isinstance(cursor, bool)
        or not isinstance(limit, int)
        or isinstance(limit, bool)
        or not 0 <= cursor <= MAX_CURSOR
        or not 1 <= limit <= MAX_EXTRACT_LIMIT
    ):
        raise ArtifactError(
            "protocol_invalid",
            "The extraction cursor or limit is out of bounds.",
            "Use cursor >= 0 and limit 1..200.",
        )
    return {"ok": True, "action": action, **_scan(source, collect_window=True, cursor=cursor, limit=limit)}
