"""Bounded PPTX and PDF creation/editing for the Orca-owned worker."""

from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path
from typing import Any, Mapping

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from pypdf import PdfReader, PdfWriter, Transformation
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from orca_excel_artifact.inputs import inspect_input


MAX_FILE_BYTES = 64 * 1024 * 1024
MAX_SLIDES = 200
MAX_OPERATIONS = 256
_CHART_TYPES = {
    "area": XL_CHART_TYPE.AREA,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
}


def _record(value: object, name: str) -> Mapping[str, Any]:
    if not isinstance(value, Mapping):
        raise ValueError(f"{name} must be an object")
    return value


def _list(value: object, name: str, maximum: int) -> list[Any]:
    if not isinstance(value, list) or len(value) > maximum:
        raise ValueError(f"{name} must be a bounded array")
    return value


def _number(value: object, default: float) -> float:
    result = default if value is None else float(value)
    if not 0 <= result <= 100:
        raise ValueError("document coordinate is outside bounds")
    return result


def _path(value: object, *, must_exist: bool) -> Path:
    path = Path(str(value))
    if not path.is_absolute() or path.is_symlink():
        raise ValueError("worker paths must be absolute regular paths")
    if must_exist and (not path.is_file() or path.stat().st_size > MAX_FILE_BYTES):
        raise ValueError("input file is unavailable or too large")
    if not must_exist and path.exists():
        raise ValueError("worker output already exists")
    return path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            size += len(chunk)
            if size > MAX_FILE_BYTES:
                raise ValueError("document exceeds the size limit")
            digest.update(chunk)
    return digest.hexdigest()


def _verify_source(request: Mapping[str, Any], kind: str) -> Path:
    source = _path(request.get("sourcePath"), must_exist=True)
    expected = request.get("expectedSha256")
    if not isinstance(expected, str) or _sha256(source) != expected:
        raise ValueError("source document hash mismatch")
    inspected = inspect_input(source, kind, expected_sha256=expected)
    if kind == "pptx":
        slides = inspected.get("presentation", {}).get("slides", [])
        unsupported = sorted(
            {
                item
                for slide in slides
                for item in slide.get("unsupportedObjects", [])
                if isinstance(item, str)
            }
        )
        if unsupported:
            raise ValueError(
                "presentation contains objects that cannot be preserved safely: "
                + ", ".join(unsupported)
            )
    return source


def _font(run: Any, style: Mapping[str, Any]) -> None:
    if "fontSize" in style:
        run.font.size = Pt(float(style["fontSize"]))
    if "bold" in style:
        run.font.bold = bool(style["bold"])
    if "italic" in style:
        run.font.italic = bool(style["italic"])
    color = style.get("color")
    if isinstance(color, str) and len(color.removeprefix("#")) == 6:
        run.font.color.rgb = RGBColor.from_string(color.removeprefix("#"))


def _text(shape: Any, value: str, style: Mapping[str, Any]) -> None:
    frame = shape.text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = value[:32767]
    _font(run, style)


def _add_element(slide: Any, element: Mapping[str, Any], artifacts: Mapping[str, Path]) -> None:
    kind = element.get("type")
    x = Inches(_number(element.get("x"), 0.5))
    y = Inches(_number(element.get("y"), 1.0))
    width = Inches(_number(element.get("width"), 9.0))
    height = Inches(_number(element.get("height"), 1.0))
    if kind == "text":
        _text(slide.shapes.add_textbox(x, y, width, height), str(element.get("text", "")), element)
    elif kind == "shape":
        _text(
            slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, width, height),
            str(element.get("text", "")),
            element,
        )
    elif kind == "image":
        image = artifacts.get(str(element.get("artifactId", "")))
        if image is None:
            raise ValueError("presentation image artifact is unavailable")
        slide.shapes.add_picture(str(image), x, y, width, height)
    elif kind == "table":
        rows = _list(element.get("rows"), "table rows", 200)
        column_count = max((len(row) for row in rows if isinstance(row, list)), default=0)
        if not rows or not 1 <= column_count <= 50:
            raise ValueError("table dimensions are invalid")
        table = slide.shapes.add_table(len(rows), column_count, x, y, width, height).table
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(_list(row, "table row", 50)):
                table.cell(row_index, column_index).text = str(value)[:32767]
    elif kind == "chart":
        chart_type = _CHART_TYPES.get(str(element.get("chartType")))
        categories = _list(element.get("categories"), "chart categories", 1000)
        series = _list(element.get("series"), "chart series", 50)
        if chart_type is None or not categories or not series:
            raise ValueError("chart specification is invalid")
        data = CategoryChartData()
        data.categories = [str(item) for item in categories]
        for item in series:
            entry = _record(item, "chart series")
            values = _list(entry.get("values"), "chart values", 1000)
            data.add_series(str(entry.get("name", "Series")), [float(value) for value in values])
        chart = slide.shapes.add_chart(chart_type, x, y, width, height, data).chart
        if element.get("title"):
            chart.has_title = True
            chart.chart_title.text_frame.text = str(element["title"])
    else:
        raise ValueError("presentation element type is unsupported")


def _add_slide(presentation: Presentation, spec: Mapping[str, Any], artifacts: Mapping[str, Path]) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if spec.get("title"):
        _add_element(
            slide,
            {
                "type": "text",
                "text": spec["title"],
                "x": 0.5,
                "y": 0.25,
                "width": 12,
                "height": 0.6,
                "fontSize": 28,
                "bold": True,
            },
            artifacts,
        )
    for element in _list(spec.get("elements", []), "slide elements", 100):
        _add_element(slide, _record(element, "slide element"), artifacts)


def _create_pptx(request: Mapping[str, Any], output: Path, artifacts: Mapping[str, Path]) -> dict[str, Any]:
    spec = _record(request.get("documentSpec"), "documentSpec")
    slides = _list(spec.get("slides"), "slides", MAX_SLIDES)
    if not slides:
        raise ValueError("presentation requires at least one slide")
    presentation = Presentation()
    if spec.get("layout") == "widescreen":
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
    for slide in slides:
        _add_slide(presentation, _record(slide, "slide"), artifacts)
    presentation.save(output)
    Presentation(output)
    return {"slideCount": len(presentation.slides), "appliedCount": len(slides)}


def _replace_paragraph_text(paragraph: Any, old: str, new: str) -> bool:
    if paragraph.text != old:
        return False
    if paragraph.runs:
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = new
    return True


def _replace_slide_text(slide: Any, old: str, new: str) -> int:
    count = 0
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            count += sum(
                _replace_paragraph_text(paragraph, old, new)
                for paragraph in shape.text_frame.paragraphs
            )
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    count += sum(
                        _replace_paragraph_text(paragraph, old, new)
                        for paragraph in cell.text_frame.paragraphs
                    )
    return count


def _edit_pptx(request: Mapping[str, Any], output: Path, artifacts: Mapping[str, Path]) -> dict[str, Any]:
    presentation = Presentation(_verify_source(request, "pptx"))
    applied = 0
    for raw in _list(request.get("operations"), "operations", MAX_OPERATIONS):
        operation = _record(raw, "operation")
        kind = operation.get("kind")
        if kind == "replace_text":
            slide_number = int(operation.get("slide", 0))
            targets = presentation.slides if slide_number == 0 else [presentation.slides[slide_number - 1]]
            applied += sum(
                _replace_slide_text(
                    slide,
                    str(operation.get("sourceText", "")),
                    str(operation.get("text", "")),
                )
                for slide in targets
            )
        elif kind == "add_slide":
            _add_slide(presentation, _record(operation.get("slideSpec"), "slideSpec"), artifacts)
            applied += 1
        elif kind == "delete_slide":
            index = int(operation.get("slide", 0)) - 1
            slide_id = presentation.slides._sldIdLst[index]
            presentation.part.drop_rel(slide_id.rId)
            del presentation.slides._sldIdLst[index]
            applied += 1
        elif kind == "add_element":
            slide = presentation.slides[int(operation.get("slide", 0)) - 1]
            _add_element(slide, _record(operation.get("element"), "element"), artifacts)
            applied += 1
        elif kind == "set_table_cell":
            slide = presentation.slides[int(operation.get("slide", 0)) - 1]
            tables = [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]
            table = tables[int(operation.get("table", 0)) - 1]
            table.cell(
                int(operation.get("row", 0)) - 1,
                int(operation.get("column", 0)) - 1,
            ).text = str(operation.get("text", ""))
            applied += 1
        else:
            raise ValueError("presentation edit operation is unsupported")
    if not presentation.slides:
        raise ValueError("presentation cannot be empty")
    presentation.save(output)
    Presentation(output)
    return {"slideCount": len(presentation.slides), "appliedCount": applied}


def _pdf_font() -> str:
    candidates = [
        Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "malgun.ttf",
        Path("/System/Library/Fonts/AppleSDGothicNeo.ttc"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for path in candidates:
        if path.is_file():
            try:
                pdfmetrics.registerFont(TTFont("OrcaDocumentFont", path))
                return "OrcaDocumentFont"
            except Exception:
                continue
    return "Helvetica"


def _draw_text_page(target: Any, page: Mapping[str, Any], size: tuple[float, float]) -> None:
    font = _pdf_font()
    _, height = size
    y = height - 54
    if page.get("title"):
        target.setFont(font, 18)
        target.drawString(54, y, str(page["title"])[:200])
        y -= 36
    font_size = float(page.get("fontSize", 11))
    target.setFont(font, font_size)
    for logical_line in str(page.get("text", "")).splitlines():
        chunks = [logical_line[index:index + 90] for index in range(0, len(logical_line), 90)] or [""]
        for line in chunks:
            if y < 54:
                target.showPage()
                target.setFont(font, font_size)
                y = height - 54
            target.drawString(54, y, line)
            y -= float(page.get("lineHeight", 16))
    target.showPage()


def _create_pdf(request: Mapping[str, Any], output: Path) -> dict[str, Any]:
    spec = _record(request.get("documentSpec"), "documentSpec")
    pages = _list(spec.get("pages"), "pages", 1000)
    if not pages:
        raise ValueError("PDF requires at least one page")
    page_size = LETTER if spec.get("pageSize") == "letter" else A4
    document = canvas.Canvas(str(output), pagesize=page_size, pageCompression=1)
    for page in pages:
        _draw_text_page(document, _record(page, "page"), page_size)
    document.save()
    return {"pageCount": len(PdfReader(output).pages), "appliedCount": len(pages)}


def _watermark(text: str, width: float, height: float) -> Any:
    stream = io.BytesIO()
    layer = canvas.Canvas(stream, pagesize=(width, height))
    layer.setFont(_pdf_font(), 36)
    layer.setFillAlpha(0.2)
    layer.saveState()
    layer.translate(width / 2, height / 2)
    layer.rotate(35)
    layer.drawCentredString(0, 0, text[:200])
    layer.restoreState()
    layer.save()
    stream.seek(0)
    return PdfReader(stream).pages[0]


def _edit_pdf(request: Mapping[str, Any], output: Path, artifacts: Mapping[str, Path]) -> dict[str, Any]:
    pages = list(PdfReader(_verify_source(request, "pdf")).pages)
    metadata: dict[str, str] = {}
    applied = 0
    for raw in _list(request.get("operations"), "operations", MAX_OPERATIONS):
        operation = _record(raw, "operation")
        kind = operation.get("kind")
        if kind == "delete_pages":
            removed = {int(value) for value in _list(operation.get("pages"), "pages", 1000)}
            pages = [page for index, page in enumerate(pages, 1) if index not in removed]
        elif kind == "reorder_pages":
            order = [int(value) for value in _list(operation.get("pages"), "pages", 1000)]
            pages = [pages[index - 1] for index in order]
        elif kind == "rotate_pages":
            selected = {int(value) for value in _list(operation.get("pages"), "pages", 1000)}
            angle = int(operation.get("degrees", 0))
            for index, page in enumerate(pages, 1):
                if index in selected:
                    page.rotate(angle)
        elif kind == "merge_pdf":
            artifact = artifacts.get(str(operation.get("artifactId", "")))
            if artifact is None:
                raise ValueError("merged PDF artifact is unavailable")
            inspect_input(artifact, "pdf")
            pages.extend(PdfReader(artifact).pages)
        elif kind == "watermark":
            for page in pages:
                page.merge_transformed_page(
                    _watermark(
                        str(operation.get("text", "")),
                        float(page.mediabox.width),
                        float(page.mediabox.height),
                    ),
                    Transformation(),
                )
        elif kind == "metadata":
            values = _record(operation.get("values"), "metadata")
            metadata = {f"/{key.title()}": str(value) for key, value in values.items()}
        else:
            raise ValueError("PDF edit operation is unsupported")
        applied += 1
    if not pages:
        raise ValueError("PDF cannot be empty")
    writer = PdfWriter()
    for page in pages:
        writer.add_page(page)
    if metadata:
        writer.add_metadata(metadata)
    with output.open("xb") as stream:
        writer.write(stream)
        stream.flush()
        os.fsync(stream.fileno())
    PdfReader(output)
    return {"pageCount": len(pages), "appliedCount": applied}


def run_document_job(value: Mapping[str, Any]) -> dict[str, Any]:
    action = str(value.get("action"))
    output = _path(value.get("outputPath"), must_exist=False)
    output.parent.mkdir(mode=0o700, parents=True, exist_ok=True)
    artifacts = {
        str(item["artifactId"]): _path(item["path"], must_exist=True)
        for item in _list(value.get("artifacts", []), "artifacts", 16)
    }
    try:
        if action == "create_pptx":
            details = _create_pptx(value, output, artifacts)
        elif action == "edit_pptx":
            details = _edit_pptx(value, output, artifacts)
        elif action == "create_pdf":
            details = _create_pdf(value, output)
        elif action == "edit_pdf":
            details = _edit_pdf(value, output, artifacts)
        else:
            raise ValueError("document action is unsupported")
        size = output.stat().st_size
        if not 0 < size <= MAX_FILE_BYTES:
            raise ValueError("document output exceeds the size limit")
        return {
            "ok": True,
            "action": action,
            "sha256": _sha256(output),
            "sizeBytes": size,
            **details,
        }
    except Exception:
        output.unlink(missing_ok=True)
        raise
